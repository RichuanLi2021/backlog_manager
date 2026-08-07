#!/usr/bin/env python3
"""
Build and export the RELEASE -> NDP -> Backlog hierarchy.

Each backlog carries both its release ("Snx Release") and its deploy plan
("Deploy Plan" / NDP). This script maps all three levels together and writes:
  - a grouped/collapsible "tree" sheet (RELEASE > NDP > Backlog)
  - a flat "mapping" sheet (one row per backlog with Release + NDP columns)

Reuses the same workspace-root .env as the other Jira skills:
  JIRA_URL=https://jira.synnex.com
  JIRA_PAT=<personal access token>     (or JIRA_EMAIL + JIRA_API_TOKEN)

Examples:
  # By release keys (release-rooted; uses the proven "Snx Release" = REL query)
  py jira-rel-review/export_hierarchy.py REL-21631 REL-20814

  # By a backlog JQL / saved filter (releases + NDPs auto-discovered from backlogs)
  py jira-rel-review/export_hierarchy.py --jql "\"(Snx) Biz Domain\" = \"TDS-GL\""
  py jira-rel-review/export_hierarchy.py --filter 12345

  # Custom output path
  py jira-rel-review/export_hierarchy.py REL-21631 --output-file jira-rel-review/hierarchy.xlsx
"""
import argparse
import base64
import json
import os
import re
import sys
import urllib.error
import urllib.parse
import urllib.request
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
WORKSPACE_ROOT = os.path.dirname(SCRIPT_DIR)
ENV_PATHS = [
    os.path.join(WORKSPACE_ROOT, ".env"),
    os.path.join(WORKSPACE_ROOT, "AIasMe", ".env"),
    os.path.expanduser("~/.cursor/skills-cursor/.env"),
]

DEPLOY_PLAN_FIELD_FALLBACK = "customfield_18302"   # Backlog -> Deploy Plan (NDP)
NDP_ADD_FIELD_FALLBACK = "customfield_20411"       # NDP Actual Deploy Date
NO_NDP = "(no NDP)"

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")


def load_dotenv(path):
    if not os.path.isfile(path):
        return
    with open(path, "r", encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            k, _, v = line.partition("=")
            k, v = k.strip(), v.strip()
            if v.startswith('"') and v.endswith('"'):
                v = v[1:-1].replace('\\"', '"')
            elif v.startswith("'") and v.endswith("'"):
                v = v[1:-1].replace("\\'", "'")
            os.environ[k] = v


def load_jira_env():
    for path in ENV_PATHS:
        load_dotenv(path)


def build_auth_header():
    pat = os.getenv("JIRA_PAT", "").strip()
    email = os.getenv("JIRA_EMAIL", "").strip()
    token = os.getenv("JIRA_API_TOKEN", "").strip()
    if pat:
        return ("Authorization", f"Bearer {pat}")
    if email and token:
        b64 = base64.b64encode(f"{email}:{token}".encode()).decode()
        return ("Authorization", f"Basic {b64}")
    return None


def request_json(url, auth_header):
    req = urllib.request.Request(
        url,
        headers={"Accept": "application/json", auth_header[0]: auth_header[1]},
    )
    with urllib.request.urlopen(req, timeout=60) as resp:
        return json.loads(resp.read().decode("utf-8", errors="replace"))


def api_get(base_url, auth_header, path):
    last_err = None
    for ver in ("2", "3"):
        url = f"{base_url}/rest/api/{ver}/{path}"
        try:
            return request_json(url, auth_header)
        except urllib.error.HTTPError as e:
            last_err = e
            if e.code == 404 and ver == "2":
                continue
            raise
    if last_err:
        raise last_err
    raise RuntimeError(f"Failed GET {path}")


def resolve_field_ids(base_url, auth_header, wanted_names):
    try:
        fields = api_get(base_url, auth_header, "field")
    except Exception:
        return {}
    mapping = {}
    for f in fields or []:
        name = (f.get("name") or "").strip().lower()
        fid = f.get("id")
        if name and fid and name in wanted_names and name not in mapping:
            mapping[name] = fid
    return mapping


def format_date(value):
    if not value:
        return "N/A"
    m = re.match(r"(\d{4})-(\d{2})-(\d{2})", str(value).strip())
    if not m:
        return str(value)
    return f"{m.group(2)}/{m.group(3)}/{m.group(1)}"


def field_text(value):
    if value in (None, "", [], {}):
        return "N/A"
    if isinstance(value, dict):
        # Cascading select: {"value": "Biz", "child": {"value": "TDS-GL"}} -> "Biz -> TDS-GL"
        if value.get("value") and isinstance(value.get("child"), dict) and value["child"].get("value"):
            return f"{value['value']} -> {value['child']['value']}"
        for key in ("displayName", "name", "value", "key"):
            if value.get(key):
                return str(value[key])
        return "N/A"
    if isinstance(value, list):
        parts = [field_text(v) for v in value]
        parts = [p for p in parts if p != "N/A"]
        return ", ".join(parts) if parts else "N/A"
    return str(value)


def extract_issue_key(value):
    if isinstance(value, dict) and value.get("key"):
        return value["key"]
    if isinstance(value, list):
        for v in value:
            k = extract_issue_key(v)
            if k != "N/A":
                return k
    if isinstance(value, str) and re.fullmatch(r"[A-Z][A-Z0-9]+-\d+", value.strip()):
        return value.strip()
    return "N/A"


def seconds_to_hours(seconds):
    if not seconds:
        return 0.0
    return round(float(seconds) / 3600.0, 2)


def issue_sort_key(key):
    m = re.fullmatch(r"([A-Z][A-Z0-9]+)-(\d+)", key or "")
    if not m:
        return ("~" + (key or ""), 0)
    return (m.group(1), int(m.group(2)))


def search_all(base_url, auth_header, jql, fields):
    # Page until a short/empty page. Do NOT rely on the reported `total`, which
    # Jira's search can transiently under-report under load (dropping trailing results).
    issues = []
    start_at = 0
    page_size = 100
    field_q = ",".join(fields)
    while True:
        q = urllib.parse.quote(jql, safe="")
        path = f"search?jql={q}&startAt={start_at}&maxResults={page_size}&fields={field_q}"
        data = api_get(base_url, auth_header, path)
        batch = data.get("issues") or []
        issues.extend(batch)
        start_at += len(batch)
        if len(batch) < page_size:
            break
    return issues


def parse_release_keys(items):
    keys = []
    for token in items:
        for part in re.split(r"[\s,]+", str(token).strip()):
            part = part.strip().upper()
            if re.fullmatch(r"REL-\d+", part):
                keys.append(part)
    return list(dict.fromkeys(keys))


def record_from_issue(issue, forced_release, snx_release_field, deploy_plan_field,
                      backlog_source_field, extra_fields=None):
    """Build one backlog record dict from a Jira issue. extra_fields: {out_key: field_id}."""
    f = issue.get("fields") or {}
    release = forced_release
    if not release and snx_release_field:
        release = extract_issue_key(f.get(snx_release_field))
    release = release or "(no Release)"
    ndp = extract_issue_key(f.get(deploy_plan_field))
    if ndp == "N/A":
        ndp = NO_NDP
    est = seconds_to_hours(
        f.get("timeoriginalestimate")
        or ((f.get("timetracking") or {}).get("originalEstimateSeconds"))
    )
    spent = seconds_to_hours(
        f.get("timespent")
        or ((f.get("timetracking") or {}).get("timeSpentSeconds"))
    )
    rec = {
        "key": issue.get("key"),
        "summary": f.get("summary") or "N/A",
        "type": ((f.get("issuetype") or {}).get("name")) or "N/A",
        "status": ((f.get("status") or {}).get("name")) or "N/A",
        "assignee": field_text(f.get("assignee")),
        "reporter": field_text(f.get("reporter")),
        "priority": field_text(f.get("priority")),
        "backlog_source": field_text(f.get(backlog_source_field)) if backlog_source_field else "N/A",
        "est": est,
        "spent": spent,
        "release": release,
        "ndp": ndp,
    }
    for out_key, fid in (extra_fields or {}).items():
        rec[out_key] = field_text(f.get(fid))
    return rec


def collect_backlogs(base_url, auth_header, rel_keys, jql, snx_release_field, deploy_plan_field,
                     backlog_source_field, issuetype):
    """
    Returns list of backlog records:
    {key, summary, type, status, assignee, reporter, priority, backlog_source,
     est, spent, release, ndp}
    """
    base_fields = [
        "summary", "issuetype", "status", "assignee", "reporter", "priority",
        "timeoriginalestimate", "timespent", "timetracking", deploy_plan_field,
    ]
    if snx_release_field:
        base_fields.append(snx_release_field)
    if backlog_source_field:
        base_fields.append(backlog_source_field)

    tagged = []  # (forced_release_or_None, issue)

    if rel_keys:
        for rel in rel_keys:
            jql_rel = f'"Snx Release" = "{rel}"'
            if issuetype:
                jql_rel += f' AND issuetype = "{issuetype}"'
            for issue in search_all(base_url, auth_header, jql_rel, base_fields):
                tagged.append((rel, issue))
    if jql:
        jql_q = jql
        if issuetype and "issuetype" not in jql_q.lower():
            jql_q = f"({jql_q}) AND issuetype = \"{issuetype}\""
        for issue in search_all(base_url, auth_header, jql_q, base_fields):
            tagged.append((None, issue))

    records = []
    seen = set()
    for forced_rel, issue in tagged:
        key = issue.get("key")
        if not key or key in seen:
            continue
        seen.add(key)
        records.append(record_from_issue(
            issue, forced_rel, snx_release_field, deploy_plan_field, backlog_source_field,
        ))
    return records


def build_tree(records):
    """release -> ndp -> [records]"""
    tree = {}
    for r in records:
        tree.setdefault(r["release"], {}).setdefault(r["ndp"], []).append(r)
    return tree


def enrich_metadata(base_url, auth_header, tree):
    """Return (release_meta, ndp_meta) dicts."""
    release_meta = {}
    ndp_meta = {}
    for rel in tree:
        if rel.startswith("REL-"):
            try:
                issue = api_get(base_url, auth_header, f"issue/{rel}?fields=summary,status")
                rf = issue.get("fields") or {}
                release_meta[rel] = {
                    "title": rf.get("summary") or "N/A",
                    "status": ((rf.get("status") or {}).get("name")) or "N/A",
                }
            except Exception:
                release_meta[rel] = {"title": "N/A", "status": "N/A"}
        else:
            release_meta[rel] = {"title": "", "status": ""}
        for ndp in tree[rel]:
            if ndp.startswith("NDP-") and ndp not in ndp_meta:
                try:
                    issue = api_get(base_url, auth_header, f"issue/{ndp}?fields=status,{NDP_ADD_FIELD_FALLBACK}")
                    nf = issue.get("fields") or {}
                    ndp_meta[ndp] = {
                        "status": ((nf.get("status") or {}).get("name")) or "N/A",
                        "add": format_date(nf.get(NDP_ADD_FIELD_FALLBACK)),
                    }
                except Exception:
                    ndp_meta[ndp] = {"status": "N/A", "add": "N/A"}
            elif ndp not in ndp_meta:
                ndp_meta[ndp] = {"status": "", "add": ""}
    return release_meta, ndp_meta


def sorted_releases(tree):
    return sorted(tree.keys(), key=issue_sort_key)


def sorted_ndps(ndp_map):
    return sorted(ndp_map.keys(), key=issue_sort_key)


def print_tree(tree, release_meta, ndp_meta):
    for rel in sorted_releases(tree):
        rm = release_meta.get(rel, {})
        ndp_map = tree[rel]
        n_backlogs = sum(len(v) for v in ndp_map.values())
        title = f" | {rm.get('title')}" if rm.get("title") else ""
        status = f" | {rm.get('status')}" if rm.get("status") else ""
        print(f"{rel}{title}{status}  ({len(ndp_map)} NDP, {n_backlogs} backlog)")
        for ndp in sorted_ndps(ndp_map):
            nm = ndp_meta.get(ndp, {})
            extra = ""
            if nm.get("status"):
                extra = f" | {nm.get('status')} | A.D.D {nm.get('add')}"
            print(f"  {ndp}{extra}  ({len(ndp_map[ndp])} backlog)")
            for r in sorted(ndp_map[ndp], key=lambda x: issue_sort_key(x["key"])):
                print(f"    {r['key']} | {r['status']} | {r['summary']}")


def autosize(ws, max_width=80):
    for col_cells in ws.columns:
        length = 0
        col = col_cells[0].column
        for cell in col_cells:
            if cell.value is not None:
                length = max(length, len(str(cell.value)))
        ws.column_dimensions[get_column_letter(col)].width = min(max_width, max(10, length + 2))


def write_excel(out_path, tree, release_meta, ndp_meta):
    wb = Workbook()

    # ---- Sheet 1: tree (grouped / collapsible) ----
    ws = wb.active
    ws.title = "tree"
    ws.sheet_properties.outlinePr.summaryBelow = False
    rel_fill = PatternFill("solid", fgColor="1F4E78")
    ndp_fill = PatternFill("solid", fgColor="2E75B6")
    hdr_fill = PatternFill("solid", fgColor="4472C4")
    white_bold = Font(bold=True, color="FFFFFF")

    tree_headers = ["Level", "Key", "Title / Summary", "Status", "Deploy Date",
                    "Assignee", "Priority", "Backlog Source"]
    ws.append(tree_headers)
    for cell in ws[1]:
        cell.fill = hdr_fill
        cell.font = white_bold

    for rel in sorted_releases(tree):
        rm = release_meta.get(rel, {})
        ndp_map = tree[rel]
        n_backlogs = sum(len(v) for v in ndp_map.values())
        ws.append([
            "RELEASE", rel, rm.get("title", ""), rm.get("status", ""), "",
            f"{len(ndp_map)} NDP", f"{n_backlogs} backlog", "",
        ])
        rel_row = ws.max_row
        for cell in ws[rel_row]:
            cell.fill = rel_fill
            cell.font = white_bold

        for ndp in sorted_ndps(ndp_map):
            nm = ndp_meta.get(ndp, {})
            ws.append([
                "NDP", ndp, "", nm.get("status", ""), nm.get("add", ""),
                f"{len(ndp_map[ndp])} backlog", "", "",
            ])
            ndp_row = ws.max_row
            ws.row_dimensions[ndp_row].outline_level = 1
            for cell in ws[ndp_row]:
                cell.fill = ndp_fill
                cell.font = Font(bold=True, color="FFFFFF")

            for r in sorted(ndp_map[ndp], key=lambda x: issue_sort_key(x["key"])):
                ws.append([
                    "Backlog", r["key"], r["summary"], r["status"], "",
                    r["assignee"], r["priority"], r["backlog_source"],
                ])
                ws.row_dimensions[ws.max_row].outline_level = 2

    ws.freeze_panes = "A2"
    autosize(ws)

    # ---- Sheet 2: flat mapping ----
    ws2 = wb.create_sheet("mapping")
    flat_headers = [
        "Release", "Release Title", "NDP", "NDP Status", "NDP Deploy Date",
        "Backlog Key", "Summary", "Issue Type", "Status", "Assignee",
        "Reporter", "Priority", "Backlog Source", "Est h", "Spent h",
    ]
    ws2.append(flat_headers)
    for cell in ws2[1]:
        cell.fill = hdr_fill
        cell.font = white_bold
    for rel in sorted_releases(tree):
        rm = release_meta.get(rel, {})
        ndp_map = tree[rel]
        for ndp in sorted_ndps(ndp_map):
            nm = ndp_meta.get(ndp, {})
            for r in sorted(ndp_map[ndp], key=lambda x: issue_sort_key(x["key"])):
                ws2.append([
                    rel, rm.get("title", ""), ndp, nm.get("status", ""), nm.get("add", ""),
                    r["key"], r["summary"], r["type"], r["status"], r["assignee"],
                    r["reporter"], r["priority"], r["backlog_source"], r["est"], r["spent"],
                ])
    ws2.freeze_panes = "A2"
    ws2.auto_filter.ref = f"A1:{get_column_letter(len(flat_headers))}{ws2.max_row}"
    autosize(ws2)

    wb.save(out_path)


def main():
    parser = argparse.ArgumentParser(description="Export RELEASE -> NDP -> Backlog hierarchy.")
    parser.add_argument("items", nargs="*", help="Release keys (REL-####), comma/space/newline separated")
    parser.add_argument("--jql", help="Backlog JQL (releases/NDPs auto-discovered from backlogs)")
    parser.add_argument("--filter", dest="filter_id", help="Saved filter/board id, becomes `filter = <id>`")
    parser.add_argument("--include-issuetype", default="Backlog",
                        help="Restrict to this issue type (default: Backlog). Use '' for all types.")
    parser.add_argument("--output-file", help="Output .xlsx path (default: auto-named in jira-rel-review/)")
    args = parser.parse_args()

    rel_keys = parse_release_keys(args.items)
    jql = args.jql
    if args.filter_id:
        jql = f"filter = {args.filter_id}"
    if not rel_keys and not jql:
        print("Usage: provide REL keys and/or --jql/--filter. Example:\n"
              "  py jira-rel-review/export_hierarchy.py REL-21631\n"
              "  py jira-rel-review/export_hierarchy.py --filter 12345", file=sys.stderr)
        sys.exit(1)

    load_jira_env()
    base_url = (os.getenv("JIRA_URL") or os.getenv("JIRA_SERVER_URL") or "").strip().rstrip("/")
    auth_header = build_auth_header()
    if not base_url or not auth_header:
        print("FAILED: set JIRA_URL and JIRA_PAT (or JIRA_EMAIL + JIRA_API_TOKEN) in workspace .env",
              file=sys.stderr)
        sys.exit(1)

    resolved = resolve_field_ids(base_url, auth_header, {"backlog source", "deploy plan", "snx release"})
    deploy_plan_field = resolved.get("deploy plan", DEPLOY_PLAN_FIELD_FALLBACK)
    backlog_source_field = resolved.get("backlog source")
    snx_release_field = resolved.get("snx release")

    issuetype = args.include_issuetype
    try:
        records = collect_backlogs(
            base_url, auth_header, rel_keys, jql,
            snx_release_field, deploy_plan_field, backlog_source_field, issuetype,
        )
    except urllib.error.HTTPError as e:
        detail = ""
        try:
            detail = e.read().decode("utf-8", errors="replace")
        except Exception:
            pass
        print(f"FAILED: search error {e.code} {e.reason} {detail}", file=sys.stderr)
        sys.exit(1)

    print(f"Collected {len(records)} backlog(s).", file=sys.stderr)
    tree = build_tree(records)
    release_meta, ndp_meta = enrich_metadata(base_url, auth_header, tree)

    print_tree(tree, release_meta, ndp_meta)

    out_path = args.output_file
    if not out_path:
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        tag = rel_keys[0] if len(rel_keys) == 1 else ("multi" if rel_keys else "jql")
        out_path = os.path.join(SCRIPT_DIR, f"Skill_jira-rel-review_hierarchy_{tag}_as_{stamp}.xlsx")
    write_excel(out_path, tree, release_meta, ndp_meta)
    print(f"\nWrote hierarchy ({len(tree)} release, {len(records)} backlog) to: {out_path}")


if __name__ == "__main__":
    main()
