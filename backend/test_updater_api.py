"""End-to-end check of the spreadsheet updater against a running API.

Exports the real report, ages it the way a week of Jira movement would — statuses
rolled back, an NDP re-opened, one backlog deleted so it looks newly raised, edits
and a highlight added on top — then uploads it and confirms the refreshed workbook
comes back current on the Jira columns and untouched everywhere else.
"""
import json
import os
import sys
import urllib.error
import urllib.request
import uuid
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import PatternFill

sys.path.insert(0, str(Path(__file__).resolve().parent))
from app.config import ADMIN_KEY  # noqa: E402

# Set BACKLOG_API when the app is served somewhere other than the default port.
API = os.getenv("BACKLOG_API", "http://127.0.0.1:8000").rstrip("/") + "/api"
KIND, QUARTER, SHEET = "GL", "2026Q3", "NA_GL"
USER = {"email": "updater.tester@tdsynnex.com", "name": "Updater Tester", "role": "SE"}
HIGHLIGHT = PatternFill("solid", fgColor="FFFF00")

failures: list[str] = []


def check(condition, message):
    if not condition:
        failures.append(message)
    return condition


def call(method, path, body=None, admin=False, raw=False):
    headers = {"Content-Type": "application/json"}
    if admin:
        headers["X-Admin-Key"] = ADMIN_KEY
    req = urllib.request.Request(
        f"{API}{path}",
        method=method,
        data=json.dumps(body).encode() if body is not None else None,
        headers=headers,
    )
    try:
        with urllib.request.urlopen(req, timeout=120) as r:
            payload = r.read()
            return payload if raw else json.loads(payload.decode())
    except urllib.error.HTTPError as e:
        return {"__error__": e.code, "detail": json.loads(e.read().decode()).get("detail")}


def post_file(path, filename, blob, fields, raw=False):
    """Multipart by hand, so the test needs nothing outside the standard library."""
    boundary = uuid.uuid4().hex
    parts = [
        f'--{boundary}\r\nContent-Disposition: form-data; name="{k}"\r\n\r\n{v}\r\n'.encode()
        for k, v in fields.items()
    ]
    parts.append(
        f'--{boundary}\r\nContent-Disposition: form-data; name="file"; '
        f'filename="{filename}"\r\nContent-Type: application/octet-stream\r\n\r\n'.encode()
    )
    parts.append(blob)
    parts.append(f"\r\n--{boundary}--\r\n".encode())
    req = urllib.request.Request(
        f"{API}{path}",
        method="POST",
        data=b"".join(parts),
        headers={"Content-Type": f"multipart/form-data; boundary={boundary}"},
    )
    try:
        with urllib.request.urlopen(req, timeout=180) as r:
            payload = r.read()
            return payload if raw else json.loads(payload.decode())
    except urllib.error.HTTPError as e:
        return {"__error__": e.code, "detail": json.loads(e.read().decode()).get("detail")}


def unlike(current, options):
    """A plausible value that is definitely not the one Jira currently holds."""
    for option in options:
        if str(option).strip().casefold() != str(current or "").strip().casefold():
            return option
    raise AssertionError(f"no alternative to {current!r} in {options}")


def age_the_file(blob, live):
    """Roll three rows backwards and drop one, so the export looks a week old.

    The old values are derived from what Jira says now rather than hard-coded, or a
    re-sync could hand us a row that already reads the way we meant to age it.
    """
    wb = load_workbook(BytesIO(blob))
    ws = wb[SHEET]
    head = {c.value: c.column for c in ws[1]}
    keys = {ws.cell(row=r, column=head["Backlog Key"]).value: r for r in range(2, ws.max_row + 1)}
    picked = list(keys)[:4]

    aged = {}
    ws.cell(row=keys[picked[0]], column=head["Status"]).value = unlike(
        live[picked[0]]["status"], ["10-Placeholder", "20-RequirementCompleted"]
    )
    aged[picked[0]] = ("Status", "status")
    ws.cell(row=keys[picked[1]], column=head["NDP Status"]).value = unlike(
        live[picked[1]]["ndp_status"], ["Open", "Verified", "IIC Approved"]
    )
    aged[picked[1]] = ("NDP Status", "ndp_status")
    ws.cell(row=keys[picked[2]], column=head["Est h"]).value = (
        float(live[picked[2]]["est_h"] or 0) + 999
    )
    aged[picked[2]] = ("Est h", "est_h")

    # Edits and a highlight the stakeholders added on top of the export.
    ws.cell(row=keys[picked[0]], column=head["Comments"]).value = "TPM: must ship first"
    ws.cell(row=keys[picked[0]], column=head["Developer"]).value = "Ann Lee"
    ws.cell(row=keys[picked[0]], column=head["Rank"]).value = 1
    for column in range(1, ws.max_column + 1):
        ws.cell(row=keys[picked[0]], column=column).fill = HIGHLIGHT

    # Deleting a row makes that backlog look like one raised after the export.
    ws.delete_rows(keys[picked[3]])

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue(), aged, picked[3], picked[0]


def main():
    call("POST", "/users", USER, admin=True)

    blob = call("GET", f"/reports/{KIND}/{QUARTER}/export?sheet={SHEET}", raw=True)
    if isinstance(blob, dict):
        print(f"nothing to export ({blob}); run a sync first")
        return 1

    upload = {"author_email": USER["email"]}

    # A file straight out of the platform is by definition current.
    fresh = post_file(f"/reports/{KIND}/{QUARTER}/refresh-file", "fresh.xlsx", blob, upload)
    if "__error__" in fresh:
        print(f"plan failed on a fresh export: {fresh}")
        return 1
    check(fresh["changes"] == [], f"a fresh export looked stale: {fresh['changes'][:3]}")
    check(fresh["new_rows"] == [], f"a fresh export was missing rows: {fresh['new_rows'][:3]}")
    check(fresh["sheets"][0]["resolved_to"] == SHEET, f"sheet not resolved: {fresh['sheets']}")
    print(f"fresh export: 0 stale cells, {fresh['unchanged_cells']} already current")

    live = {r["key"]: r for r in call("GET", f"/reports/{KIND}/{QUARTER}")["items"]
            if r["sheet"] == SHEET}
    aged, expected, deleted, edited = age_the_file(blob, live)
    plan = post_file(f"/reports/{KIND}/{QUARTER}/refresh-file", "week_old.xlsx", aged, upload)
    if "__error__" in plan:
        print(f"plan failed: {plan}")
        return 1
    found = {(c["key"], c["label"]) for c in plan["changes"]}
    print(f"aged file: {len(plan['changes'])} stale cells across {plan['field_counts']}, "
          f"{len(plan['new_rows'])} raised since, {len(plan['gone_rows'])} dropped")

    for key, (label, _) in expected.items():
        check((key, label) in found, f"{key} {label} was not detected as stale")
    check([n["key"] for n in plan["new_rows"]] == [deleted],
          f"the deleted row should look newly raised: {plan['new_rows']}")

    updated = post_file(
        f"/reports/{KIND}/{QUARTER}/refresh-file/download", "week_old.xlsx", aged,
        {**upload, "fields": "", "add_new": "true", "flag_gone": "true"}, raw=True,
    )
    if isinstance(updated, dict):
        print(f"download failed: {updated}")
        return 1
    print(f"downloaded {len(updated)} bytes")

    ws = load_workbook(BytesIO(updated))[SHEET]
    head = {c.value: c.column for c in ws[1]}
    rows = {ws.cell(row=r, column=head["Backlog Key"]).value: r for r in range(2, ws.max_row + 1)}

    # Excel gives whole floats back as ints and blanks as None, so compare on meaning.
    def alike(cell, platform):
        if isinstance(cell, (int, float)) or isinstance(platform, (int, float)):
            return float(cell or 0) == float(platform or 0)
        return str(cell or "").strip() == str(platform or "").strip()

    for key, (label, name) in expected.items():
        got = ws.cell(row=rows[key], column=head[label]).value
        want = live[key][name]
        check(alike(got, want), f"{key} {label} is {got!r}, Jira says {want!r}")

    check(ws.cell(row=rows[edited], column=head["Comments"]).value == "TPM: must ship first",
          "the comment was overwritten")
    check(ws.cell(row=rows[edited], column=head["Developer"]).value == "Ann Lee",
          "the developer was overwritten")
    check(ws.cell(row=rows[edited], column=head["Rank"]).value == 1, "the rank was overwritten")
    check(ws.cell(row=rows[edited], column=head["Comments"]).fill.fgColor.rgb == "00FFFF00",
          "the highlight did not survive")
    print(f"{edited}: Jira columns current, comment / developer / rank / highlight intact")

    check(deleted in rows, "the backlog raised since the export was not appended")
    if deleted in rows:
        marker = next((c.column for c in ws[1] if c.value == "Update Status"), None)
        check(marker is not None, "no Update Status column was added")
        if marker:
            check("new since" in (ws.cell(row=rows[deleted], column=marker).value or ""),
                  "the appended row was not marked as new")
        check(ws.cell(row=rows[deleted], column=head["Summary"]).value ==
              live[deleted]["summary"], "the appended row has the wrong summary")

    log = load_workbook(BytesIO(updated))["update log"]
    entries = [r for r in log.iter_rows(values_only=True) if r[0] == SHEET]
    check(len(entries) == len(plan["changes"]) + len(plan["new_rows"]),
          f"log has {len(entries)} rows for {len(plan['changes'])} changes "
          f"and {len(plan['new_rows'])} additions")
    print(f"update log: {len(entries)} entries")

    # Refreshing the refreshed file must find nothing left to do.
    again = post_file(f"/reports/{KIND}/{QUARTER}/refresh-file", "again.xlsx", updated, upload)
    check(again.get("changes") == [], f"a refreshed file still looks stale: {again.get('changes', [])[:3]}")
    check(again.get("new_rows") == [], f"rows still missing after a refresh: {again.get('new_rows')}")
    print(f"re-checked the output: {len(again.get('changes', []))} stale cells left")

    denied = post_file(
        f"/reports/{KIND}/{QUARTER}/refresh-file", "week_old.xlsx", aged,
        {"author_email": "nobody@example.com"},
    )
    check(denied.get("__error__") == 404, f"a non-member should be refused, got {denied}")
    junk = post_file(f"/reports/{KIND}/{QUARTER}/refresh-file", "notes.txt", b"hello", upload)
    check(junk.get("__error__") == 400, f"a non-workbook should be refused, got {junk}")
    print(f"non-member -> {denied.get('__error__')}, non-workbook -> {junk.get('__error__')}")

    for u in call("GET", "/users"):
        if u["email"] == USER["email"]:
            call("DELETE", f"/users/{u['id']}", admin=True)

    if failures:
        print("\nFAILURES:")
        for f in failures:
            print(f"  - {f}")
        return 1
    print("\nPASS: a week-old export comes back current on Jira and unchanged elsewhere")
    return 0


if __name__ == "__main__":
    sys.exit(main())
