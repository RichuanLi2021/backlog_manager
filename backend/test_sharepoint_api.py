"""End-to-end: refresh a week-old file in a SharePoint folder, and age the NDPs.

Self-contained: starts its own API on a free port, against a copy of the database
and a temp folder standing in for the OneDrive-synced GL_BL folder. Nothing it does
reaches the real database, the real SharePoint, or a server already running.

Needs Jira credentials, because the NDP ages are read from the live changelog.
"""
import json
import os
import shutil
import socket
import subprocess
import sys
import tempfile
import time
import urllib.error
import urllib.request
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
from app.config import ADMIN_KEY  # noqa: E402

KIND, QUARTER = "GL", "2026Q4"
USER = {"email": "sharepoint.tester@tdsynnex.com", "name": "SharePoint Tester", "role": "SE"}
failures: list[str] = []


def check(condition, message):
    if not condition:
        failures.append(message)


def free_port() -> int:
    with socket.socket() as s:
        s.bind(("127.0.0.1", 0))
        return s.getsockname()[1]


class Api:
    def __init__(self, port):
        self.base = f"http://127.0.0.1:{port}/api"

    def call(self, method, path, body=None, admin=False, raw=False):
        headers = {"Content-Type": "application/json"}
        if admin:
            headers["X-Admin-Key"] = ADMIN_KEY
        req = urllib.request.Request(
            f"{self.base}{path}",
            method=method,
            headers=headers,
            data=json.dumps(body).encode() if body is not None else None,
        )
        try:
            with urllib.request.urlopen(req, timeout=120) as r:
                payload = r.read()
                return payload if raw else json.loads(payload.decode())
        except urllib.error.HTTPError as e:
            return {"__error__": e.code, "detail": json.loads(e.read().decode()).get("detail")}


def age_a_week(blob: bytes) -> tuple[bytes, dict]:
    """Roll statuses back, change an NDP, and drop a row, as a week-old copy would be."""
    wb = load_workbook(BytesIO(blob))
    ws = wb[wb.sheetnames[0]]
    head = {c.value: c.column for c in ws[1]}
    rows = list(range(2, min(ws.max_row, 8) + 1))

    def key_of(r):
        return ws.cell(row=r, column=head["Backlog Key"]).value

    picked = {
        "status": key_of(rows[0]),
        "ndp_status": key_of(rows[1]),
        "dropped": key_of(rows[2]),
        "sheet": ws.title,
    }
    status_now = ws.cell(row=rows[0], column=head["Status"]).value
    ws.cell(row=rows[0], column=head["Status"]).value = (
        "10-Placeholder" if status_now != "10-Placeholder" else "20-RequirementCompleted"
    )
    ndp_now = ws.cell(row=rows[1], column=head["NDP Status"]).value
    ws.cell(row=rows[1], column=head["NDP Status"]).value = (
        "Verified" if ndp_now != "Verified" else "Open"
    )
    ws.cell(row=rows[0], column=head["Comments"]).value = "TPM: keep this comment"
    ws.delete_rows(rows[2])
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue(), picked


def main():
    work = Path(tempfile.mkdtemp(prefix="sp_api_"))
    shutil.copy(HERE / "backlog.db", work / "copy.db")
    folder = work / "GL_BL" / QUARTER / "Updated"
    folder.mkdir(parents=True)

    port = free_port()
    env = {
        **os.environ,
        "DATABASE_URL": f"sqlite:///{work / 'copy.db'}",
        "SHAREPOINT_DIR": str(work / "GL_BL"),
        "SHAREPOINT_BACKUP_DIR": str(work / "backups"),
    }
    server = subprocess.Popen(
        [sys.executable, "-m", "uvicorn", "app.main:app", "--port", str(port)],
        cwd=HERE,
        env=env,
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
    )
    api = Api(port)
    try:
        for _ in range(60):
            try:
                if api.call("GET", "/health").get("status") == "ok":
                    break
            except OSError:
                time.sleep(0.5)
        else:
            print("server did not start")
            return 1
        return run(api, folder)
    finally:
        server.terminate()
        server.wait(timeout=10)
        shutil.rmtree(work, ignore_errors=True)


def run(api: Api, folder: Path) -> int:
    api.call("POST", "/users", USER, admin=True)

    # ---- an empty folder, then the team's working file ---------------------------
    empty = api.call("GET", f"/sharepoint/{KIND}/{QUARTER}")
    check(empty["found"] and empty["files"] == [], f"empty folder misreported: {empty}")

    blob = api.call("GET", f"/reports/{KIND}/{QUARTER}/export", raw=True)
    if isinstance(blob, dict):
        print(f"cannot export {KIND} {QUARTER}: {blob}")
        return 1
    aged, picked = age_a_week(blob)
    target = folder / f"GL_Summary_{QUARTER}.xlsx"
    target.write_bytes(aged)
    other = f"AR_Summary_{QUARTER}.xlsx"
    (folder / other).write_bytes(blob)
    (folder / f"~$GL_Summary_{QUARTER}_old.xlsx").write_bytes(b"lock")

    listing = api.call("GET", f"/sharepoint/{KIND}/{QUARTER}")
    names = {f["name"]: f for f in listing["files"]}
    check(set(names) == {target.name, other}, f"listing should skip the lock file: {sorted(names)}")
    check(names[target.name]["matches_report"], "the GL file should match the GL report")
    check(not names[other]["matches_report"], "an AR file matched the GL report")
    print(f"folder: {sorted(names)} (lock file skipped)")

    # ---- scan: where the file stands against Jira now -----------------------------
    scan = api.call("POST", f"/sharepoint/{KIND}/{QUARTER}/scan", {"name": target.name})
    if "__error__" in scan:
        print(f"scan failed: {scan}")
        return 1
    plan = scan["plan"]
    moved = {(c["key"], c["field"]) for c in plan["changes"]}
    check((picked["status"], "status") in moved, f"rolled-back status not found: {picked}")
    check((picked["ndp_status"], "ndp_status") in moved, "rolled-back NDP status not found")
    check(
        picked["dropped"] in {n["key"] for n in plan["new_rows"]},
        f"the dropped row should come back as raised since: {plan['new_rows'][:3]}",
    )
    progress = plan["progress"]
    check(progress is not None, "scan should carry the progress summary")
    if progress:
        # The export carries archived rows too, so the file can hold backlogs Jira has
        # dropped. The counts must reconcile through exactly those and the raised ones.
        expected_now = (
            progress["backlogs_in_file"] - len(plan["gone_rows"]) + len(plan["new_rows"])
        )
        check(
            progress["backlogs_now"] == expected_now,
            f"progress does not reconcile: {progress['backlogs_in_file']} in file - "
            f"{len(plan['gone_rows'])} dropped + {len(plan['new_rows'])} raised = {expected_now}, "
            f"but Jira now shows {progress['backlogs_now']}",
        )
        print(
            f"scan: {len(plan['changes'])} stale cells, {len(plan['new_rows'])} raised, {len(plan['gone_rows'])} dropped; "
            f"progress {progress['backlogs_in_file']} -> {progress['backlogs_now']} backlogs "
            f"across {len(progress['backlogs'])} statuses, {len(progress['ndps'])} NDP statuses"
        )

    body = {
        "name": target.name,
        "token": scan["file"]["token"],
        "author_email": USER["email"],
        "fields": [],
        "add_new": True,
        "flag_gone": True,
    }

    # ---- the guards ----------------------------------------------------------------
    stranger = api.call(
        "POST", f"/sharepoint/{KIND}/{QUARTER}/update", {**body, "author_email": "nobody@example.com"}
    )
    check(stranger.get("__error__") == 404, f"a non-member should be refused: {stranger}")

    escape = api.call("POST", f"/sharepoint/{KIND}/{QUARTER}/scan", {"name": "../../copy.db"})
    check(escape.get("__error__") == 400, f"a path outside the folder should be refused: {escape}")

    before = target.read_bytes()
    time.sleep(0.05)
    target.write_bytes(before)  # a colleague saves: same content, new timestamp
    stale = api.call("POST", f"/sharepoint/{KIND}/{QUARTER}/update", body)
    check(
        stale.get("__error__") == 409 and "changed after you previewed" in (stale.get("detail") or ""),
        f"an update after someone else saved should be refused: {stale}",
    )
    check(target.read_bytes() == before, "a refused update still changed the file")
    print(
        f"guards: non-member {stranger.get('__error__')}, outside path {escape.get('__error__')}, "
        f"stale preview {stale.get('__error__')}"
    )

    # ---- the update ------------------------------------------------------------------
    rescan = api.call("POST", f"/sharepoint/{KIND}/{QUARTER}/scan", {"name": target.name})
    result = api.call(
        "POST", f"/sharepoint/{KIND}/{QUARTER}/update", {**body, "token": rescan["file"]["token"]}
    )
    if "__error__" in result:
        print(f"update failed: {result}")
        return 1
    check(
        result["cells_updated"] == len(rescan["plan"]["changes"]),
        f"cells updated {result['cells_updated']} != planned {len(rescan['plan']['changes'])}",
    )
    check(Path(result["backup"]).read_bytes() == before, "the backup is not the pre-update file")

    sheet = load_workbook(BytesIO(target.read_bytes()))[picked["sheet"]]
    head = {c.value: c.column for c in sheet[1]}
    rows = {
        sheet.cell(row=r, column=head["Backlog Key"]).value: r for r in range(2, sheet.max_row + 1)
    }
    check(
        sheet.cell(row=rows[picked["status"]], column=head["Comments"]).value
        == "TPM: keep this comment",
        "the update overwrote a comment",
    )
    check(picked["dropped"] in rows, "the backlog raised since was not appended in SharePoint")

    after = api.call("POST", f"/sharepoint/{KIND}/{QUARTER}/scan", {"name": target.name})
    check(
        after["plan"]["changes"] == [] and after["plan"]["new_rows"] == [],
        f"after the update the file should be current: {len(after['plan']['changes'])} stale",
    )
    print(
        f"update: {result['cells_updated']} cells, {result['rows_added']} added, comment kept; "
        f"re-scan finds {len(after['plan']['changes'])} stale cells"
    )

    # ---- NDP ages, from the live Jira changelog ---------------------------------------
    refreshed = api.call("POST", f"/reports/{KIND}/{QUARTER}/ndp-ages")
    check(
        "__error__" not in refreshed and refreshed.get("from_jira", 0) > 0,
        f"NDP ages should come from Jira: {refreshed}",
    )
    rules = api.call("GET", "/ndp-age-rules")
    tree = api.call("GET", f"/reports/{KIND}/{QUARTER}/tree")
    ndps = {n["ndp"]: n for s in tree for r in s["releases"] for n in r["ndps"]}
    levels: dict[str, int] = {}
    for n in ndps.values():
        levels[n["age_level"]] = levels.get(n["age_level"], 0) + 1
        if n["age_level"] in ("green", "yellow", "red"):
            days = n["age_days"]
            if days <= rules["warn_days"]:
                want = "green"
            elif days <= rules["flag_days"]:
                want = "yellow"
            else:
                want = "red"
            check(n["age_level"] == want, f"{n['ndp']} at {days}d is {n['age_level']}, want {want}")
        if n["status"] in rules["done_statuses"]:
            check(n["age_level"] == "done", f"{n['ndp']} is {n['status']} but aged {n['age_level']}")
    print(f"NDP ages: {refreshed.get('from_jira')} read from Jira; levels {levels}")

    if failures:
        print("\nFAILURES:")
        for f in failures:
            print(f"  - {f}")
        return 1
    print("\nPASS: a week-old SharePoint file is scanned, guarded, updated in place, and NDPs age")
    return 0


if __name__ == "__main__":
    sys.exit(main())
