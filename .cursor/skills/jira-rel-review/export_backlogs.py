#!/usr/bin/env python3
"""
Export a backlog-centric list from a Jira JQL / saved board filter and associate
each backlog with its Deploy Plan (NDP).

Reuses the same workspace-root .env as jira-rel-review / jira-ndp-review / jira-backlog-spec:
  JIRA_URL=https://jira.synnex.com
  JIRA_PAT=<personal access token>     (or JIRA_EMAIL + JIRA_API_TOKEN)

Examples:
  # By a saved board/filter id
  py jira-rel-review/export_backlogs.py --filter 12345

  # By raw JQL (quote it)
  py jira-rel-review/export_backlogs.py --jql "\"(Snx) Biz Domain\" = \"TDS-GL\" AND issuetype = Backlog"

  # Custom output path (default: auto-named .xlsx in jira-rel-review/)
  py jira-rel-review/export_backlogs.py --filter 12345 --output-file jira-rel-review/backlogs.xlsx
"""
import argparse
import base64
import json
import os
import re
import sys
import urllib.error
import urllib.parse
import urllib.request
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
WORKSPACE_ROOT = os.path.dirname(SCRIPT_DIR)
ENV_PATHS = [
    os.path.join(WORKSPACE_ROOT, ".env"),
    os.path.join(WORKSPACE_ROOT, "AIasMe", ".env"),
]

# Known Jira field ids (fallbacks if name resolution fails).
DEPLOY_PLAN_FIELD_FALLBACK = "customfield_18302"   # Backlog -> Deploy Plan (NDP) link
NDP_ADD_FIELD_FALLBACK = "customfield_20411"       # NDP Actual Deploy Date

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
if hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")


def load_dotenv(path):
    if not os.path.isfile(path):
        return
    with open(path, "r", encoding="utf-8") as f:
        for line in f:
            line = line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            k, _, v = line.partition("=")
            k, v = k.strip(), v.strip()
            if v.startswith('"') and v.endswith('"'):
                v = v[1:-1].replace('\\"', '"')
            elif v.startswith("'") and v.endswith("'"):
                v = v[1:-1].replace("\\'", "'")
            os.environ[k] = v


def load_jira_env():
    for path in ENV_PATHS:
        load_dotenv(path)


def build_auth_header():
    pat = os.getenv("JIRA_PAT", "").strip()
    email = os.getenv("JIRA_EMAIL", "").strip()
    token = os.getenv("JIRA_API_TOKEN", "").strip()
    if pat:
        return ("Authorization", f"Bearer {pat}")
    if email and token:
        b64 = base64.b64encode(f"{email}:{token}".encode()).decode()
        return ("Authorization", f"Basic {b64}")
    return None


def request_json(url, auth_header, method="GET"):
    req = urllib.request.Request(
        url,
        headers={"Accept": "application/json", auth_header[0]: auth_header[1]},
        method=method,
    )
    with urllib.request.urlopen(req, timeout=60) as resp:
        return json.loads(resp.read().decode("utf-8", errors="replace"))


def api_get(base_url, auth_header, path):
    """GET a Jira REST path, trying api/2 then api/3."""
    last_err = None
    for ver in ("2", "3"):
        url = f"{base_url}/rest/api/{ver}/{path}"
        try:
            return request_json(url, auth_header)
        except urllib.error.HTTPError as e:
            last_err = e
            if e.code == 404 and ver == "2":
                continue
            raise
    if last_err:
        raise last_err
    raise RuntimeError(f"Failed GET {path}")


def resolve_field_ids(base_url, auth_header, wanted_names):
    """Return {display_name_lower: field_id} for the requested field display names."""
    try:
        fields = api_get(base_url, auth_header, "field")
    except Exception:
        return {}
    mapping = {}
    for f in fields or []:
        name = (f.get("name") or "").strip().lower()
        fid = f.get("id")
        if name and fid and name in wanted_names and name not in mapping:
            mapping[name] = fid
    return mapping


def format_date(value):
    if not value:
        return "N/A"
    m = re.match(r"(\d{4})-(\d{2})-(\d{2})", str(value).strip())
    if not m:
        return str(value)
    return f"{m.group(2)}/{m.group(3)}/{m.group(1)}"


def field_text(value):
    if value in (None, "", [], {}):
        return "N/A"
    if isinstance(value, dict):
        # Cascading select: {"value": "Biz", "child": {"value": "TDS-GL"}} -> "Biz -> TDS-GL"
        if value.get("value") and isinstance(value.get("child"), dict) and value["child"].get("value"):
            return f"{value['value']} -> {value['child']['value']}"
        for key in ("displayName", "name", "value", "key"):
            if value.get(key):
                return str(value[key])
        return "N/A"
    if isinstance(value, list):
        parts = [field_text(v) for v in value]
        parts = [p for p in parts if p != "N/A"]
        return ", ".join(parts) if parts else "N/A"
    return str(value)


def deploy_plan_key(value):
    """Deploy Plan field is usually an issue link object with a .key."""
    if isinstance(value, dict) and value.get("key"):
        return value["key"]
    if isinstance(value, list):
        for v in value:
            k = deploy_plan_key(v)
            if k != "N/A":
                return k
    if isinstance(value, str) and re.fullmatch(r"[A-Z][A-Z0-9]+-\d+", value.strip()):
        return value.strip()
    return "N/A"


def search_all(base_url, auth_header, jql, fields):
    """Paginate the Jira search endpoint and return all issues."""
    issues = []
    start_at = 0
    page_size = 100
    field_q = ",".join(fields)
    while True:
        q = urllib.parse.quote(jql, safe="")
        path = f"search?jql={q}&startAt={start_at}&maxResults={page_size}&fields={field_q}"
        data = api_get(base_url, auth_header, path)
        batch = data.get("issues") or []
        issues.extend(batch)
        total = data.get("total", len(issues))
        start_at += len(batch)
        if not batch or start_at >= total:
            break
    return issues


def autosize(ws, max_width=70):
    for col_cells in ws.columns:
        length = 0
        col = col_cells[0].column
        for cell in col_cells:
            val = cell.value
            if val is not None:
                length = max(length, len(str(val)))
        ws.column_dimensions[get_column_letter(col)].width = min(max_width, max(10, length + 2))


def main():
    parser = argparse.ArgumentParser(description="Export backlogs with associated NDP to Excel.")
    src = parser.add_mutually_exclusive_group(required=True)
    src.add_argument("--jql", help="Raw JQL query")
    src.add_argument("--filter", dest="filter_id", help="Saved filter/board id, becomes `filter = <id>`")
    parser.add_argument(
        "--include-issuetype",
        default="Backlog",
        help="Restrict to this issue type (default: Backlog). Use '' to include all types.",
    )
    parser.add_argument("--output-file", help="Output .xlsx path (default: auto-named in jira-rel-review/)")
    parser.add_argument("--no-ndp-details", action="store_true", help="Skip fetching NDP status / deploy date")
    args = parser.parse_args()

    load_jira_env()
    base_url = (os.getenv("JIRA_URL") or os.getenv("JIRA_SERVER_URL") or "").strip().rstrip("/")
    auth_header = build_auth_header()
    if not base_url or not auth_header:
        print("FAILED: set JIRA_URL and JIRA_PAT (or JIRA_EMAIL + JIRA_API_TOKEN) in workspace .env", file=sys.stderr)
        sys.exit(1)

    if args.filter_id:
        jql = f"filter = {args.filter_id}"
    else:
        jql = args.jql
    if args.include_issuetype:
        # Only append if the caller didn't already constrain issuetype.
        if "issuetype" not in jql.lower():
            jql = f"({jql}) AND issuetype = \"{args.include_issuetype}\""

    # Resolve custom fields by display name (robust across instances).
    resolved = resolve_field_ids(
        base_url,
        auth_header,
        {"backlog source", "deploy plan"},
    )
    backlog_source_field = resolved.get("backlog source")
    deploy_plan_field = resolved.get("deploy plan", DEPLOY_PLAN_FIELD_FALLBACK)

    request_fields = ["summary", "issuetype", "status", "assignee", "reporter", "priority", deploy_plan_field]
    if backlog_source_field:
        request_fields.append(backlog_source_field)

    print(f"JQL: {jql}", file=sys.stderr)
    try:
        issues = search_all(base_url, auth_header, jql, request_fields)
    except urllib.error.HTTPError as e:
        detail = ""
        try:
            detail = e.read().decode("utf-8", errors="replace")
        except Exception:
            pass
        print(f"FAILED: search error {e.code} {e.reason} {detail}", file=sys.stderr)
        sys.exit(1)
    print(f"Fetched {len(issues)} issue(s).", file=sys.stderr)

    # Optionally enrich with NDP status + actual deploy date (fetch each NDP once).
    ndp_cache = {}

    def get_ndp_details(ndp_key):
        if ndp_key == "N/A" or args.no_ndp_details:
            return ("N/A", "N/A")
        if ndp_key in ndp_cache:
            return ndp_cache[ndp_key]
        try:
            ndp = api_get(base_url, auth_header, f"issue/{ndp_key}?fields=status,{NDP_ADD_FIELD_FALLBACK}")
            nf = ndp.get("fields") or {}
            status = ((nf.get("status") or {}).get("name")) or "N/A"
            add = format_date(nf.get(NDP_ADD_FIELD_FALLBACK))
        except Exception:
            status, add = "N/A", "N/A"
        ndp_cache[ndp_key] = (status, add)
        return ndp_cache[ndp_key]

    headers = [
        "Issue Type", "Issue Key", "Summary", "Status", "Assignee",
        "Reporter", "Priority", "Backlog Source", "NDP", "NDP Status", "NDP Deploy Date",
    ]
    rows = []
    for issue in issues:
        f = issue.get("fields") or {}
        ndp_key = deploy_plan_key(f.get(deploy_plan_field))
        ndp_status, ndp_add = get_ndp_details(ndp_key)
        backlog_source = field_text(f.get(backlog_source_field)) if backlog_source_field else "N/A"
        rows.append([
            ((f.get("issuetype") or {}).get("name")) or "N/A",
            issue.get("key") or "N/A",
            f.get("summary") or "N/A",
            ((f.get("status") or {}).get("name")) or "N/A",
            field_text(f.get("assignee")),
            field_text(f.get("reporter")),
            field_text(f.get("priority")),
            backlog_source,
            ndp_key,
            ndp_status,
            ndp_add,
        ])

    # Sort by NDP then issue key for grouping readability.
    def sort_key(r):
        m = re.fullmatch(r"([A-Z][A-Z0-9]+)-(\d+)", r[8] or "")
        ndp_sort = (m.group(1), int(m.group(2))) if m else ("~", 0)
        mk = re.fullmatch(r"([A-Z][A-Z0-9]+)-(\d+)", r[1] or "")
        key_sort = (mk.group(1), int(mk.group(2))) if mk else (r[1], 0)
        return (ndp_sort, key_sort)

    rows.sort(key=sort_key)

    wb = Workbook()
    ws = wb.active
    ws.title = "backlogs"
    header_fill = PatternFill("solid", fgColor="4472C4")
    header_font = Font(bold=True, color="FFFFFF")
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(vertical="center")
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
    autosize(ws)

    out_path = args.output_file
    if not out_path:
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        out_path = os.path.join(SCRIPT_DIR, f"Skill_jira-rel-review_backlogs_as_{stamp}.xlsx")
    wb.save(out_path)
    print(f"Wrote {len(rows)} backlog rows to: {out_path}")


if __name__ == "__main__":
    main()
