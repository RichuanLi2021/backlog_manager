"""Compare synced DB rows against the skill's existing Excel workbooks, per sheet."""
import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import func, select

sys.path.insert(0, str(Path(__file__).resolve().parent))
from app.db import SessionLocal  # noqa: E402
from app.models import BacklogSnapshot  # noqa: E402

QUARTER = sys.argv[1] if len(sys.argv) > 1 else "2026Q3"
EXCEL_ROOT = Path.home() / ".cursor" / "skills" / "jira-rel-review" / "Report"


def excel_counts(kind: str) -> Counter:
    path = EXCEL_ROOT / kind / f"{kind}_Summary_{QUARTER}.xlsx"
    counts: Counter = Counter()
    if not path.is_file():
        return counts
    wb = load_workbook(path, read_only=True, data_only=True)
    for ws in wb.worksheets:
        if ws.title == "summary" or ws.title.endswith("_tree"):
            continue
        header = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
        if "Backlog Key" not in header:
            continue
        col = header.index("Backlog Key")
        n = sum(
            1
            for row in ws.iter_rows(min_row=2, values_only=True)
            if row[col] not in (None, "")
        )
        counts[ws.title] = n
    wb.close()
    return counts


def keys_for(db, kind: str, sheet: str) -> set:
    return set(
        db.scalars(
            select(BacklogSnapshot.key).where(
                BacklogSnapshot.report_kind == kind,
                BacklogSnapshot.quarter == QUARTER,
                BacklogSnapshot.sheet == sheet,
                BacklogSnapshot.archived.is_(False),
            )
        )
    )


def check_internal_consistency(db):
    """Excel files are snapshots from different moments; these invariants are not."""
    print("=== internal consistency ===")
    ok = True
    for region in ("NA", "LAC"):
        combined = keys_for(db, "AR_GL", f"{region}_AR+GL")
        union = keys_for(db, "AR_GL", f"{region}_AR") | keys_for(db, "AR_GL", f"{region}_GL")
        flag = "OK  " if combined == union else "DIFF"
        ok &= combined == union
        print(f"  {flag} {region}_AR+GL ({len(combined)}) == {region}_AR + {region}_GL ({len(union)})")

    for kind, sheets in (("AR", ("NA_AR", "LAC_AR")), ("GL", ("NA_GL", "LAC_GL"))):
        for sheet in sheets:
            a, b = keys_for(db, kind, sheet), keys_for(db, "AR_GL", sheet)
            flag = "OK  " if a == b else "DIFF"
            ok &= a == b
            print(f"  {flag} {kind}.{sheet} ({len(a)}) == AR_GL.{sheet} ({len(b)})")

    print(f"  -> {'all invariants hold' if ok else 'INCONSISTENT'}\n")


def main():
    db = SessionLocal()
    print(f"Quarter {QUARTER}   (excel = last skill run, may predate this sync)\n")
    for kind in ("AR_GL", "AR", "GL"):
        db_counts = Counter(
            dict(
                db.execute(
                    select(BacklogSnapshot.sheet, func.count(BacklogSnapshot.id))
                    .where(
                        BacklogSnapshot.report_kind == kind,
                        BacklogSnapshot.quarter == QUARTER,
                        BacklogSnapshot.archived.is_(False),
                    )
                    .group_by(BacklogSnapshot.sheet)
                ).all()
            )
        )
        xl_counts = excel_counts(kind)
        print(f"=== {kind} ===")
        for sheet in sorted(set(db_counts) | set(xl_counts)):
            d, x = db_counts.get(sheet, 0), xl_counts.get(sheet, 0)
            flag = "OK  " if d == x else "DIFF"
            print(f"  {flag} {sheet:10} db={d:4}  excel={x:4}")
        print(f"  total db={sum(db_counts.values())}  excel={sum(xl_counts.values())}\n")
    check_internal_consistency(db)
    db.close()


if __name__ == "__main__":
    main()
