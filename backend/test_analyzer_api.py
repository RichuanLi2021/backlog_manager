"""End-to-end check of the spreadsheet analyzer against a running API.

Exports the report the platform produces, rebuilds it the way a colleague sends one
back — a title row on top, renamed headers, only some columns, their own ranking —
uploads it, applies part of the preview, and confirms the saved view, the untouched
live view, and the group check. `test_analyzer.py` covers the parsing rules without
a server; this covers the two endpoints and what they write.
"""
import json
import os
import sys
import urllib.error
import urllib.request
import uuid
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent))
from app.config import ADMIN_KEY  # noqa: E402
from app.services.excel_export import priority_tab  # noqa: E402

# Set BACKLOG_API when the app is served somewhere other than the default port.
API = os.getenv("BACKLOG_API", "http://127.0.0.1:8000").rstrip("/") + "/api"
KIND, QUARTER, SHEET = "GL", "2026Q3", "NA_GL"
# Exports name each tab as the team's working copy, e.g. Q3_NA_GL_Priority_List.
TAB = priority_tab(SHEET, QUARTER)
REPORT_KINDS = ("GL", "AR", "AR_GL")
USER = {"email": "analyzer.tester@tdsynnex.com", "name": "Analyzer Tester", "role": "TPM"}


def call(method, path, body=None, admin=False, raw=False):
    headers = {"Content-Type": "application/json"}
    if admin:
        headers["X-Admin-Key"] = ADMIN_KEY
    req = urllib.request.Request(
        f"{API}{path}",
        method=method,
        data=json.dumps(body).encode() if body is not None else None,
        headers=headers,
    )
    try:
        with urllib.request.urlopen(req, timeout=120) as r:
            payload = r.read()
            return payload if raw else json.loads(payload.decode())
    except urllib.error.HTTPError as e:
        return {"__error__": e.code, "detail": json.loads(e.read().decode()).get("detail")}


def post_file(path, filename, blob, fields, raw=False):
    """Multipart by hand, so the test needs nothing outside the standard library."""
    boundary = uuid.uuid4().hex
    parts = [
        f'--{boundary}\r\nContent-Disposition: form-data; name="{k}"\r\n\r\n{v}\r\n'.encode()
        for k, v in fields.items()
    ]
    parts.append(
        f'--{boundary}\r\nContent-Disposition: form-data; name="file"; '
        f'filename="{filename}"\r\nContent-Type: application/octet-stream\r\n\r\n'.encode()
    )
    parts.append(blob)
    parts.append(f"\r\n--{boundary}--\r\n".encode())
    req = urllib.request.Request(
        f"{API}{path}",
        method="POST",
        data=b"".join(parts),
        headers={"Content-Type": f"multipart/form-data; boundary={boundary}"},
    )
    try:
        with urllib.request.urlopen(req, timeout=120) as r:
            payload = r.read()
            return payload if raw else json.loads(payload.decode())
    except urllib.error.HTTPError as e:
        return {"__error__": e.code, "detail": json.loads(e.read().decode()).get("detail")}


def colleagues_copy(headers, rows, key_col):
    """The same backlogs, reversed ranking, renamed columns, plus a row we do not have."""
    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet(SHEET)
    ws.append(["GL backlog - TPM review, please merge"])
    ws.append([])
    ws.append(["#", "Backlog", "Summary", "Dev Owner", "UAT", "Go-Live", "Remarks"])
    picked = rows[:6]
    for i, row in enumerate(picked):
        ws.append([
            len(picked) - i,
            row[key_col],
            row[headers.index("Summary")],
            "Imported Dev" if i < 2 else "",
            "2026-09-30" if i == 0 else "",
            "",
            f"TPM note {i}" if i < 3 else "",
        ])
    ws.append([99, "CISGL-999999", "not in the report", "", "", "", "ghost row"])
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue(), [r[key_col] for r in picked]


def every_sheet_is_covered(failures):
    """A whole-report export must be analyzed tab by tab, not just the first one.

    The Export button used to always scope to the visible sheet, so uploads only ever
    carried one tab and only that tab could be compared. This walks every report,
    including AR_GL with its six sheets and the `+` in their names.
    """
    for kind in REPORT_KINDS:
        detail = call("GET", f"/reports/{kind}/{QUARTER}")
        expected = detail.get("sheets") or []
        if not expected:
            print(f"{kind}: no rows yet, skipped")
            continue
        blob = call("GET", f"/reports/{kind}/{QUARTER}/export", raw=True)
        if isinstance(blob, dict):
            failures.append(f"{kind}: whole-report export failed: {blob}")
            continue

        result = post_file(
            f"/reports/{kind}/{QUARTER}/analyze", f"{kind}_all.xlsx", blob,
            {"author_email": USER["email"], "target_role": "TPM", "compare_with": "LIVE"},
        )
        if "__error__" in result:
            failures.append(f"{kind}: analyze rejected the whole-report export: {result}")
            continue

        covered = {s["name"]: s for s in result["sheet_summary"]}
        # Tabs are exported as each region's priority list; each must read as exactly
        # its own region sheet, never a neighbour that shares its backlogs.
        for sheet in expected:
            tab = priority_tab(sheet, QUARTER)
            got = covered.get(tab, {}).get("target_sheets")
            if got != [sheet]:
                failures.append(f"{kind}: tab {tab} should read as [{sheet}], got {got}")
        reached = {t for s in covered.values() for t in s["target_sheets"]}
        missing = [s for s in expected if s not in reached]
        if missing:
            failures.append(f"{kind}: these sheets were never read: {missing}")
        unmatched = [
            f"{s['name']} {s['matched']}/{s['rows_read']}"
            for s in covered.values()
            if s["rows_read"] and s["matched"] != s["rows_read"]
        ]
        if unmatched:
            failures.append(f"{kind}: rows read but not matched: {unmatched}")
        print(f"{kind}: {len(covered)} tabs read, "
              + ", ".join(f"{s['name']} {s['matched']}" for s in covered.values()))


def write_back_settles_the_file(failures, keys):
    """A kept change stays in the file; a discarded one goes back to the view's value.

    The point of writing back is that the reviewer's copy and the view they saved end
    up saying the same thing, so the discarded half matters as much as the kept half.
    """
    blob = call("GET", f"/reports/{KIND}/{QUARTER}/export?sheet={SHEET}", raw=True)
    source = load_workbook(BytesIO(blob))[TAB]
    headers = [c.value for c in source[1]]
    key_col, comment_col = headers.index("Backlog Key"), headers.index("IT Comment")

    wb = load_workbook(BytesIO(blob))
    ws = wb[TAB]
    rows = {ws.cell(row=r, column=key_col + 1).value: r for r in range(2, ws.max_row + 1)}
    keep, drop = keys[0], keys[1]
    ws.cell(row=rows[keep], column=comment_col + 1).value = "keep this one"
    ws.cell(row=rows[drop], column=comment_col + 1).value = "discard this one"
    buffer = BytesIO()
    wb.save(buffer)
    edited = buffer.getvalue()

    upload = {"author_email": USER["email"], "target_role": "SE", "compare_with": "LIVE"}
    result = post_file(f"/reports/{KIND}/{QUARTER}/analyze", "wb.xlsx", edited, upload)
    if "__error__" in result:
        failures.append(f"write-back analyze failed: {result}")
        return

    # Keep the first backlog's comment, leave the second one out of the selection.
    kept = [
        {"sheet": r["sheet"], "key": r["key"], "field": c["field"], "value": c["incoming"]}
        for r in result["rows"]
        for c in r["changes"]
        if r["key"] == keep and c["field"] == "comments"
    ]
    if not kept:
        failures.append("the comment meant to be kept was not offered as a change")
        return

    body = {**upload, "changes": json.dumps(kept)}
    plan = post_file(f"/reports/{KIND}/{QUARTER}/analyze/file-preview", "wb.xlsx", edited, body)
    if "__error__" in plan:
        failures.append(f"file preview failed: {plan}")
        return
    touched = {(c["key"], c["field"]) for c in plan["changes"]}
    check(failures, (drop, "comments") in touched,
          f"the discarded comment should be reverted in the file: {sorted(touched)[:5]}")
    check(failures, (keep, "comments") not in touched,
          "the kept comment should not be rewritten, it is already right")
    print(f"file preview: {len(plan['changes'])} cells would change, "
          f"discarded {drop} reverts, kept {keep} stays")

    written = post_file(
        f"/reports/{KIND}/{QUARTER}/analyze/file-download", "wb.xlsx", edited, body, raw=True
    )
    if isinstance(written, dict):
        failures.append(f"file download failed: {written}")
        return
    out = load_workbook(BytesIO(written))[TAB]
    back = {out.cell(row=r, column=key_col + 1).value: r for r in range(2, out.max_row + 1)}
    check(failures,
          out.cell(row=back[keep], column=comment_col + 1).value == "keep this one",
          "the kept comment did not survive into the downloaded file")
    check(failures,
          out.cell(row=back[drop], column=comment_col + 1).value != "discard this one",
          "the discarded comment was left in the downloaded file")
    check(failures, len(back) == len(rows),
          f"the file gained or lost rows: {len(back)} vs {len(rows)}")
    print(f"file download: {len(written)} bytes, {len(back)} rows, kept and discarded "
          "both settled")


def check(failures, condition, message):
    if not condition:
        failures.append(message)


def main():
    failures = []
    call("POST", "/users", USER, admin=True)
    every_sheet_is_covered(failures)

    blob = call("GET", f"/reports/{KIND}/{QUARTER}/export?sheet={SHEET}", raw=True)
    if isinstance(blob, dict):
        print(f"nothing to export ({blob}); run a sync first")
        return 1
    source = load_workbook(BytesIO(blob))[TAB]
    headers = [c.value for c in source[1]]
    rows = [[c.value for c in r] for r in source.iter_rows(min_row=2)]
    if len(rows) < 6:
        print(f"{SHEET} has only {len(rows)} rows; run a sync first")
        return 1
    edited, keys = colleagues_copy(headers, rows, headers.index("Backlog Key"))
    print(f"built a renamed copy of {len(keys)} rows from a {len(blob)}-byte export")

    upload = {"author_email": USER["email"], "target_role": "TPM"}
    result = post_file(f"/reports/{KIND}/{QUARTER}/analyze", "tpm_review.xlsx", edited, upload)
    if "__error__" in result:
        print(f"analyze failed: {result}")
        return 1
    print(f"analyze: {result['rows_read']} rows read, {result['keys_matched']} matched, "
          f"{result['unchanged_cells']} already identical")
    print(f"columns recognized: {result['detected_columns']}")

    if result["unmatched_keys"] != ["CISGL-999999"]:
        failures.append(f"key outside the report not reported: {result['unmatched_keys']}")
    if result["keys_matched"] != len(keys):
        failures.append(f"expected {len(keys)} matches, got {result['keys_matched']}")

    changes = [
        {"sheet": r["sheet"], "key": r["key"], "field": c["field"], "value": c["incoming"]}
        for r in result["rows"]
        for c in r["changes"]
    ]
    fields = sorted({c["field"] for c in changes})
    print(f"{len(changes)} cell changes across {fields}")
    if not {"rank", "comments", "developer"} <= set(fields):
        failures.append(f"expected rank, comments and developer changes, got {fields}")

    # Leave the developer column behind, because a preview people cannot trim is not
    # a preview. What is dropped here has to stay dropped in the saved view.
    kept = [c for c in changes if c["field"] != "developer"]
    snap = call("POST", f"/reports/{KIND}/{QUARTER}/analyze/apply", {
        "author_email": USER["email"],
        "target_role": "TPM",
        "label": "Imported from tpm_review.xlsx",
        "based_on_id": (result.get("baseline") or {}).get("id"),
        "changes": kept,
    })
    if "__error__" in snap:
        print(f"apply failed: {snap}")
        return 1
    print(f"applied {len(kept)}/{len(changes)} changes as snapshot {snap['id']} "
          f"({snap['role']}, {snap['row_count']} rows)")

    view = call("GET", f"/reports/{KIND}/{QUARTER}?snapshot_id={snap['id']}")
    saved = {r["key"]: r for r in view["items"] if r["sheet"] == SHEET}
    if saved[keys[-1]]["rank"] != 1:
        failures.append(f"ranking not imported: {keys[-1]} is rank {saved[keys[-1]]['rank']}")
    if saved[keys[0]]["comments"] != "TPM note 0":
        failures.append(f"comment not imported: {saved[keys[0]]['comments']!r}")
    if saved[keys[0]]["developer"]:
        failures.append("developer was applied although the preview excluded it")
    print(f"saved view: {keys[-1]} is rank 1, {keys[0]} reads '{saved[keys[0]]['comments']}', "
          f"developer left empty")

    again = post_file(
        f"/reports/{KIND}/{QUARTER}/analyze", "tpm_review.xlsx", edited,
        {**upload, "compare_with": "TPM"},
    )
    remaining = sorted({c["field"] for r in again["rows"] for c in r["changes"]})
    print(f"vs TPM View: {again['unchanged_cells']} cells identical, "
          f"{len(again['identical'])} rows fully agree, still differs on {remaining}")
    if remaining != ["developer"]:
        failures.append(f"applied changes came back as differences: {remaining}")
    if not again["identical"]:
        failures.append("rows that now agree were not reported as identical")
    if "TPM View" not in again["baseline_label"]:
        failures.append(f"baseline not named: {again['baseline_label']}")

    # The same file against the live report has to show everything as incoming again,
    # because the live view carries none of the text a snapshot holds.
    live_cmp = post_file(
        f"/reports/{KIND}/{QUARTER}/analyze", "tpm_review.xlsx", edited,
        {**upload, "compare_with": "LIVE"},
    )
    live_fields = sorted({c["field"] for r in live_cmp["rows"] for c in r["changes"]})
    print(f"vs live report: {len(live_cmp['rows'])} rows differ across {live_fields}, "
          f"+{live_cmp['added_cells']} added ~{live_cmp['changed_cells']} changed, "
          f"{live_cmp['stale_jira_cells']} stale Jira cells, {live_cmp['elapsed_ms']} ms")
    if not {"comments", "developer"} <= set(live_fields):
        failures.append(f"live comparison lost the file's text: {live_fields}")
    if len(live_cmp["rows"]) <= len(again["rows"]):
        failures.append(
            f"live baseline should differ more than the TPM one: "
            f"{len(live_cmp['rows'])} vs {len(again['rows'])}"
        )
    if live_cmp["added_cells"] + live_cmp["changed_cells"] != sum(
        len(r["changes"]) for r in live_cmp["rows"]
    ):
        failures.append("added/changed counts do not add up to the listed changes")
    if "Latest backlog report" not in live_cmp["baseline_label"]:
        failures.append(f"live baseline not named: {live_cmp['baseline_label']}")

    bad = post_file(
        f"/reports/{KIND}/{QUARTER}/analyze", "tpm_review.xlsx", edited,
        {**upload, "compare_with": "NOPE"},
    )
    if bad.get("__error__") != 400:
        failures.append(f"an unknown baseline should be refused, got {bad}")

    write_back_settles_the_file(failures, keys)

    live = call("GET", f"/reports/{KIND}/{QUARTER}")
    if any(r["developer"] or r["comments"] for r in live["items"]):
        failures.append("the import leaked into the live Jira view")

    denied = post_file(
        f"/reports/{KIND}/{QUARTER}/analyze", "tpm_review.xlsx", edited,
        {"author_email": "nobody@example.com", "target_role": "TPM"},
    )
    if denied.get("__error__") != 404:
        failures.append(f"a non-member should be refused, got {denied}")
    junk = post_file(f"/reports/{KIND}/{QUARTER}/analyze", "notes.txt", b"hello", upload)
    if junk.get("__error__") != 400:
        failures.append(f"a non-workbook should be refused, got {junk}")
    print(f"non-member -> {denied.get('__error__')}, non-workbook -> {junk.get('__error__')}")

    # Leave the database as it was found.
    call("DELETE", f"/snapshots/{snap['id']}")
    for u in call("GET", "/users"):
        if u["email"] == USER["email"]:
            call("DELETE", f"/users/{u['id']}", admin=True)
    print("cleaned up the test snapshot and user")

    if failures:
        print("\nFAILURES:")
        for f in failures:
            print(f"  - {f}")
        return 1
    print("\nPASS: a renamed copy previews, applies in part, and never touches the live view")
    return 0


if __name__ == "__main__":
    sys.exit(main())
