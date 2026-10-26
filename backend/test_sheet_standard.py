"""Checks for the standard-layout pass on working tabs.

The tab below copies the real Q4_NA_GL_Priority_List header: Package wedged after
Assignee, the ETAs before the hours, one Comments column, Update Status last. The
pass must land it in the standard layout without losing a value, a fill or a width,
leave the frozen NA_GL tab alone, and do nothing the second time.
"""
import sys
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

sys.path.insert(0, str(Path(__file__).resolve().parent))
from app.services.excel_export import COLUMNS  # noqa: E402
from app.services.sheet_standard import describe, standardize  # noqa: E402

failures: list[str] = []
YELLOW = PatternFill("solid", fgColor="FFFF00")

PRIORITY_HEADER = [
    "Rank", "Release", "Release Title", "NDP", "NDP Status", "NDP Deploy Date",
    "Backlog Key", "Summary", "Issue Type", "Status", "Assignee", "package", "Developer",
    "UAT ETA", "PROD ETA", "Est h", "Spent h", "Comments", "Reporter", "Priority",
    "Biz Domain", "Req Qtr", "Labels", "Backlog Source", "Risk", "Update Status",
]
OLD_NA_GL = [
    "Rank", "Release", "Release Title", "NDP", "NDP Status", "NDP Deploy Date",
    "Backlog Key", "Summary", "Issue Type", "Status", "Assignee", "Developer", "Reporter",
    "Priority", "Biz Domain", "Req Qtr", "Labels", "Backlog Source", "Est h", "Spent h",
    "UAT ETA", "PROD ETA", "Comments",
]
STANDARD = [title for _, title in COLUMNS]


def check(condition, message):
    if not condition:
        failures.append(message)


def row(header, **values):
    return [values.get(h, "") for h in header]


def working_file() -> bytes:
    wb = Workbook()
    na = wb.active
    na.title = "NA_GL"
    na.append(OLD_NA_GL)
    na.append(row(OLD_NA_GL, **{"Backlog Key": "CISGL-1", "Comments": "frozen note"}))

    ws = wb.create_sheet("Q4_NA_GL_Priority_List")
    ws.append(PRIORITY_HEADER)
    ws.append(row(PRIORITY_HEADER, **{
        "Rank": 1, "Backlog Key": "CISGL-1", "Status": "70-InCoding", "package": "gl-portal",
        "Developer": "divya", "UAT ETA": "2026-10-23", "Est h": 40,
        "Comments": "DEV will start soon after CISGL-2208 is deployed",
        "Risk": "high", "Update Status": "new since this file was exported",
    }))
    ws.append(row(PRIORITY_HEADER, **{"Rank": 2, "Backlog Key": "CISGL-2", "Comments": "Need confirmation"}))
    for cell in ws[2]:
        cell.fill = YELLOW  # someone highlighted the first backlog by hand
    ws.column_dimensions["R"].width = 61  # Comments, made wide by hand
    ws.auto_filter.ref = "A1:Z3"
    # A filter someone set on UAT ETA (column N here) and on Status (column J).
    ws.auto_filter.add_filter_column(PRIORITY_HEADER.index("UAT ETA"), ["2026-10-23"])
    ws.auto_filter.add_filter_column(PRIORITY_HEADER.index("Status"), ["70-InCoding"])
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def test_standard_layout():
    data, notes = standardize(working_file(), "GL.xlsx")
    wb = load_workbook(BytesIO(data))
    ws = wb["Q4_NA_GL_Priority_List"]
    header = [c.value for c in ws[1] if c.value is not None]
    check(header == STANDARD + ["Risk", "Update Status"],
          f"header should be the standard, then kept extras, then Update Status:\n  {header}")

    col = {c.value: c.column for c in ws[1] if c.value}
    first = {h: ws.cell(row=2, column=c).value for h, c in col.items()}
    check(first["IT Comment"] == "DEV will start soon after CISGL-2208 is deployed",
          f"the old Comments text should be in IT Comment: {first['IT Comment']!r}")
    check(first["TPM Comment"] is None, "TPM Comment should start empty")
    check(first["Package"] == "gl-portal", f"package value lost: {first['Package']!r}")
    check(first["Developer"] == "divya" and first["UAT ETA"] == "2026-10-23",
          "developer or UAT ETA did not move with their column")
    check(first["Est h"] == 40, f"hours did not move with their column: {first['Est h']!r}")
    check(first["Risk"] == "high", "an unrecognised column must be kept, not dropped")
    check(first["Update Status"] == "new since this file was exported", "Update Status lost")
    check(col["Package"] == col["Assignee"] + 1, "Package should sit right after Assignee")
    check(col["TPM Comment"] == col["IT Comment"] + 1, "TPM Comment should follow IT Comment")
    check(ws.cell(row=3, column=col["IT Comment"]).value == "Need confirmation",
          "second row's comment did not move")

    fills = {h: ws.cell(row=2, column=c).fill.fgColor.rgb for h, c in col.items()}
    check(fills["IT Comment"] == "00FFFF00", "the highlight did not move with the comment")
    check(fills["TPM Comment"] == "00FFFF00", "the new column should carry the row's highlight")
    from openpyxl.utils import get_column_letter
    check(ws.column_dimensions[get_column_letter(col["IT Comment"])].width == 61,
          "the hand-set width did not move with the comment column")
    check(ws.auto_filter.ref.startswith("A1:"), f"filter range lost: {ws.auto_filter.ref}")
    on = {header[f.colId]: list(f.filters.filter) for f in ws.auto_filter.filterColumn}
    check(on == {"UAT ETA": ["2026-10-23"], "Status": ["70-InCoding"]},
          f"filters must stay on the columns they were set on, wherever those moved: {on}")

    frozen = [c.value for c in wb["NA_GL"][1] if c.value]
    check(frozen == OLD_NA_GL, "the frozen NA_GL tab must be left exactly as it was")
    check(len(notes) == 1 and notes[0].startswith("Q4_NA_GL_Priority_List:"),
          f"one note for the one tab rearranged: {notes}")
    print(f"layout: {notes[0]}")
    return data


def test_second_pass_changes_nothing(data):
    again, notes = standardize(data, "GL.xlsx")
    check(notes == [] and again == data, f"a standard tab must be left alone: {notes}")
    check(describe(data, "GL.xlsx") == [], "describe should report nothing to do")
    print("second pass: nothing to do, bytes unchanged")


def test_a_file_without_priority_tabs_is_standardised_whole():
    """Uploads and the AR workbook may have only region tabs: those are the working ones."""
    wb = Workbook()
    ws = wb.active
    ws.title = "NA_AR"
    ws.append(OLD_NA_GL)
    ws.append(row(OLD_NA_GL, **{"Backlog Key": "CISAR-1", "Comments": "ar note"}))
    buffer = BytesIO()
    wb.save(buffer)
    data, notes = standardize(buffer.getvalue(), "AR.xlsx")
    out = load_workbook(BytesIO(data))["NA_AR"]
    header = [c.value for c in out[1] if c.value]
    check(header == STANDARD, f"NA_AR should be in the standard layout: {header}")
    check(out.cell(row=2, column=header.index("IT Comment") + 1).value == "ar note",
          "the AR comment did not land in IT Comment")
    print(f"region-only file: {notes[0]}")


def main():
    data = None
    try:
        data = test_standard_layout()
    except Exception as exc:
        failures.append(f"test_standard_layout raised {type(exc).__name__}: {exc}")
    for test, args in ((test_second_pass_changes_nothing, (data,)),
                       (test_a_file_without_priority_tabs_is_standardised_whole, ())):
        try:
            if data is not None or not args:
                test(*args)
        except Exception as exc:
            failures.append(f"{test.__name__} raised {type(exc).__name__}: {exc}")
    if failures:
        print("\nFAILURES:")
        for f in failures:
            print(f"  - {f}")
        return 1
    print("\nPASS: working tabs land in the standard layout with nothing lost, once")
    return 0


if __name__ == "__main__":
    sys.exit(main())
