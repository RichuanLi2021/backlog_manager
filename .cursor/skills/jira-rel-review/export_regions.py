#!/usr/bin/env python3
"""
Fetch and organize RELEASE -> NDP -> Backlog for two regions (NA / LAC) across
biz-domain combinations (TDS-AR, TDS-GL, and both together), filtered by quarter.

Config: jira-rel-review/regions.config.json  (field ids/values, editable).
Env:    workspace-root .env  (JIRA_URL + JIRA_PAT).

Latest-report workflow (fixed paths, update in place):
  py jira-rel-review/update_latest_reports.py --quarter 2026Q3

Writes three workbooks under Report/AR_GL, Report/AR, Report/GL.
"""
import argparse
import json
import os
import sys
from datetime import datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from export_hierarchy import (  # noqa: E402
    NO_NDP,
    build_tree,
    enrich_metadata,
    issue_sort_key,
    load_jira_env,
    record_from_issue,
    resolve_field_ids,
    search_all,
    sorted_ndps,
    sorted_releases,
    build_auth_header,
)

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
CONFIG_PATH = os.path.join(SCRIPT_DIR, "regions.config.json")

# Columns filled from Jira on every refresh
MAPPING_HEADERS = [
    "Release", "Release Title", "NDP", "NDP Status", "NDP Deploy Date",
    "Backlog Key", "Summary", "Issue Type", "Status", "Assignee", "Developer",
    "Reporter", "Priority", "Biz Domain", "Req Qtr", "Labels", "Backlog Source",
    "Est h", "Spent h", "UAT ETA", "PROD ETA", "Comments",
]
# User-editable columns preserved across updates (keyed by Backlog Key)
MANUAL_COLS = ["Developer", "UAT ETA", "PROD ETA", "Comments"]

# Backlog exclusion / region classification rules
EXCLUDE_KEY_PREFIXES = ("KSPE",)
EXCLUDE_LABEL_TOKENS = ("gl-killsp",)
LAC_KEY_PREFIXES = ("MCAM", "WMCC")
NA_KEY_PREFIXES = ("CISGL", "CISAR", "CISAR-", "CISGL-")
HDR_FILL = PatternFill("solid", fgColor="4472C4")
REL_FILL = PatternFill("solid", fgColor="1F4E78")
NDP_FILL = PatternFill("solid", fgColor="2E75B6")
WHITE_BOLD = Font(bold=True, color="FFFFFF")
TS_LABEL = Font(bold=True, size=12)


def load_config():
    with open(CONFIG_PATH, "r", encoding="utf-8") as f:
        return json.load(f)


def report_output_path(kind, quarter):
    """Fixed report path: Report/<kind>/<kind>_Summary_<quarter>.xlsx"""
    folder = os.path.join(SCRIPT_DIR, "Report", kind)
    os.makedirs(folder, exist_ok=True)
    return os.path.join(folder, f"{kind}_Summary_{quarter}.xlsx")


def report_kind(domain_labels, cfg):
    atoms = set()
    for label in domain_labels:
        for v in cfg["domains"].get(label, [label]):
            atoms.add(v.replace("TDS-", "").strip().upper())
    if {"AR", "GL"}.issubset(atoms):
        return "AR_GL"
    if atoms == {"GL"}:
        return "GL"
    if atoms == {"AR"}:
        return "AR"
    return "_".join(sorted(atoms)) or "report"


def _labels_text(record):
    labels = record.get("labels", "")
    if isinstance(labels, list):
        return " ".join(str(x).lower() for x in labels)
    return str(labels).lower()


def should_exclude_record(record):
    key = (record.get("key") or "").upper()
    if any(key.startswith(p) for p in EXCLUDE_KEY_PREFIXES):
        return True
    label_text = _labels_text(record)
    return any(tok in label_text for tok in EXCLUDE_LABEL_TOKENS)


def infer_region(record):
    """Classify backlog as NA or LAC by key prefix / labels. Returns None if excluded."""
    if should_exclude_record(record):
        return None
    key = (record.get("key") or "").upper()
    # NA GL/AR keys always win (CISGL / CISAR)
    if key.startswith("CISGL") or key.startswith("CISAR"):
        return "NA"
    if any(key.startswith(p) for p in LAC_KEY_PREFIXES):
        return "LAC"
    labels = record.get("labels", "")
    if isinstance(labels, list):
        tokens = [str(x).lower().strip() for x in labels]
    else:
        tokens = [t.lower().strip() for t in str(labels).replace(",", " ").split() if t.strip()]
    if "lac" in tokens:
        return "LAC"
    if key.startswith("CIS"):
        return "NA"
    return None


def filter_records_for_region(records, region_label):
    """Keep only backlogs belonging to the target region; drop excluded keys."""
    out = []
    dropped_exclude = 0
    dropped_wrong_region = 0
    for r in records:
        if should_exclude_record(r):
            dropped_exclude += 1
            continue
        inferred = infer_region(r)
        if inferred != region_label:
            dropped_wrong_region += 1
            continue
        out.append(r)
    if dropped_exclude or dropped_wrong_region:
        print(f"  filter {region_label}: kept {len(out)}, excluded {dropped_exclude}, "
              f"wrong-region {dropped_wrong_region}", file=sys.stderr)
    return out


def jql_domain_clause(domain_field, values):
    quoted = ",".join(f'"{v}"' for v in values)
    return f'"{domain_field}" in ({quoted})'


def build_combo(base_url, auth_header, cfg, region_base_jql, domain_values,
                quarter, snx_release_field, deploy_plan_field, backlog_source_field):
    domain_field = cfg["domain_field"]
    biz_domain_field = cfg.get("biz_domain_field")
    req_qtr_field = cfg.get("req_qtr_field")

    scoped = region_base_jql.replace("{quarter}", quarter)
    jql = f'({scoped}) AND {jql_domain_clause(domain_field, domain_values)}'

    fetch_fields = [
        "summary", "issuetype", "status", "assignee", "reporter", "priority",
        "labels", "timeoriginalestimate", "timespent", "timetracking",
        deploy_plan_field, snx_release_field,
    ]
    if backlog_source_field:
        fetch_fields.append(backlog_source_field)
    if biz_domain_field:
        fetch_fields.append(biz_domain_field)
    if req_qtr_field and req_qtr_field not in fetch_fields:
        fetch_fields.append(req_qtr_field)

    issues = search_all(base_url, auth_header, jql, fetch_fields)

    extra = {}
    if biz_domain_field:
        extra["biz_domain"] = biz_domain_field
    if req_qtr_field:
        extra["req_qtr"] = req_qtr_field
    extra["labels"] = "labels"

    records = []
    seen = set()
    for issue in issues:
        key = issue.get("key")
        if not key or key in seen:
            continue
        seen.add(key)
        records.append(record_from_issue(
            issue, None, snx_release_field, deploy_plan_field, backlog_source_field, extra_fields=extra,
        ))
    return jql, records


def autosize(ws, max_width=80):
    for col_cells in ws.columns:
        length = 0
        col = col_cells[0].column
        for cell in col_cells:
            if cell.value is not None:
                length = max(length, len(str(cell.value)))
        ws.column_dimensions[get_column_letter(col)].width = min(max_width, max(10, length + 2))


def col_index(headers, name):
    return headers.index(name) + 1


def extract_manual_data(path):
    """Return {sheet_name: {backlog_key: {col: value}}} for user-editable columns."""
    if not path or not os.path.isfile(path):
        return {}
    try:
        wb = load_workbook(path, data_only=True)
    except Exception:
        return {}
    out = {}
    for ws in wb.worksheets:
        if ws.title == "summary":
            continue
        hdr = [c.value for c in ws[1]]
        if "Backlog Key" not in hdr:
            continue
        key_col = col_index(hdr, "Backlog Key")
        manual = {}
        for name in MANUAL_COLS:
            if name in hdr:
                manual[name] = col_index(hdr, name)
        if not manual:
            continue
        sheet_data = {}
        for r in range(2, ws.max_row + 1):
            bk = ws.cell(row=r, column=key_col).value
            if not bk:
                continue
            vals = {}
            for name, c in manual.items():
                v = ws.cell(row=r, column=c).value
                if v is not None and str(v).strip() != "":
                    vals[name] = v
            if vals:
                sheet_data[str(bk).strip()] = vals
        if sheet_data:
            out[ws.title] = sheet_data
    wb.close()
    return out


def mapping_row(r, rm, ndp, nm, manual=None):
  manual = manual or {}
  return [
      r["release"], rm.get("title", ""), ndp, nm.get("status", ""), nm.get("add", ""),
      r["key"], r["summary"], r["type"], r["status"], r["assignee"],
      manual.get("Developer", ""),
      r["reporter"], r["priority"], r.get("biz_domain", "N/A"), r.get("req_qtr", "N/A"),
      r.get("labels", "N/A"), r["backlog_source"], r["est"], r["spent"],
      manual.get("UAT ETA", ""), manual.get("PROD ETA", ""), manual.get("Comments", ""),
  ]


def add_mapping_sheet(wb, sheet_name, tree, release_meta, ndp_meta, manual_by_key=None):
    manual_by_key = manual_by_key or {}
    ws = wb.create_sheet(sheet_name[:31])
    ws.append(MAPPING_HEADERS)
    for cell in ws[1]:
        cell.fill = HDR_FILL
        cell.font = WHITE_BOLD
    for rel in sorted_releases(tree):
        rm = release_meta.get(rel, {})
        for ndp in sorted_ndps(tree[rel]):
            nm = ndp_meta.get(ndp, {})
            for r in sorted(tree[rel][ndp], key=lambda x: issue_sort_key(x["key"])):
                manual = manual_by_key.get(r["key"], {})
                ws.append(mapping_row(r, rm, ndp, nm, manual))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(MAPPING_HEADERS))}{ws.max_row}"
    autosize(ws)


def add_tree_sheet(wb, sheet_name, tree, release_meta, ndp_meta):
    ws = wb.create_sheet(sheet_name[:31])
    ws.sheet_properties.outlinePr.summaryBelow = False
    headers = ["Level", "Key", "Title / Summary", "Status", "Deploy Date",
               "Assignee", "Priority", "Backlog Source"]
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = HDR_FILL
        cell.font = WHITE_BOLD
    for rel in sorted_releases(tree):
        rm = release_meta.get(rel, {})
        ndp_map = tree[rel]
        n_backlogs = sum(len(v) for v in ndp_map.values())
        ws.append(["RELEASE", rel, rm.get("title", ""), rm.get("status", ""), "",
                   f"{len(ndp_map)} NDP", f"{n_backlogs} backlog", ""])
        for cell in ws[ws.max_row]:
            cell.fill = REL_FILL
            cell.font = WHITE_BOLD
        for ndp in sorted_ndps(ndp_map):
            nm = ndp_meta.get(ndp, {})
            ws.append(["NDP", ndp, "", nm.get("status", ""), nm.get("add", ""),
                       f"{len(ndp_map[ndp])} backlog", "", ""])
            ws.row_dimensions[ws.max_row].outline_level = 1
            for cell in ws[ws.max_row]:
                cell.fill = NDP_FILL
                cell.font = WHITE_BOLD
            for r in sorted(ndp_map[ndp], key=lambda x: issue_sort_key(x["key"])):
                ws.append(["Backlog", r["key"], r["summary"], r["status"], "",
                           r["assignee"], r["priority"], r["backlog_source"]])
                ws.row_dimensions[ws.max_row].outline_level = 2
    ws.freeze_panes = "A2"
    autosize(ws)


def add_summary_sheet(wb, summary_rows, updated_at, quarter):
    ws = wb.create_sheet("summary", 0)
    ts = updated_at.strftime("%m/%d/%Y %I:%M:%S %p")
    ws.append([f"Last Updated: {ts}"])
    ws["A1"].font = TS_LABEL
    ws.append([f"Quarter: {quarter}"])
    ws.append([])
    headers = ["Region", "Domain", "Quarter", "# Releases", "# NDPs", "# Backlogs", "JQL (region+domain)"]
    ws.append(headers)
    header_row = ws.max_row
    for cell in ws[header_row]:
        cell.fill = HDR_FILL
        cell.font = WHITE_BOLD
    for row in summary_rows:
        ws.append(row)
    ws.freeze_panes = f"A{header_row + 1}"
    autosize(ws, max_width=120)


def fetch_domain_pool(base_url, auth_header, cfg, quarter, domain_values,
                      snx_release_field, deploy_plan_field, backlog_source_field):
    """Fetch records for a domain from every region base JQL, dedupe, drop exclusions."""
    merged = {}
    for src_region, region_cfg in cfg["regions"].items():
        region_base_jql = region_cfg.get("base_jql") if isinstance(region_cfg, dict) else region_cfg
        if not region_base_jql:
            continue
        for dv in domain_values:
            _, recs = build_combo(
                base_url, auth_header, cfg, region_base_jql, [dv],
                quarter, snx_release_field, deploy_plan_field, backlog_source_field,
            )
            for r in recs:
                if should_exclude_record(r):
                    continue
                merged[r["key"]] = r
    return list(merged.values())


def build_report_workbook(quarter, domain_labels=None, region_labels=None):
    """Fetch Jira data and return (combo_data, summary_rows).
    combo_data: list of (sheet_name, tree, release_meta, ndp_meta).
    Records are fetched from all region JQLs then reclassified into NA/LAC sheets.
    """
    cfg = load_config()
    if domain_labels is None:
        domain_labels = list(cfg["domains"].keys())
    if region_labels is None:
        region_labels = list(cfg["regions"].keys())

    load_jira_env()
    base_url = (os.getenv("JIRA_URL") or os.getenv("JIRA_SERVER_URL") or "").strip().rstrip("/")
    auth_header = build_auth_header()
    if not base_url or not auth_header:
        raise RuntimeError("set JIRA_URL and JIRA_PAT in workspace .env")

    resolved = resolve_field_ids(base_url, auth_header, {cfg.get("backlog_source_field_name", "backlog source").lower()})
    backlog_source_field = resolved.get(cfg.get("backlog_source_field_name", "backlog source").lower())
    snx_release_field = cfg.get("snx_release_field")
    deploy_plan_field = cfg.get("deploy_plan_field")

    combo_data = []
    summary_rows = []

    for domain_label in domain_labels:
        domain_values = cfg["domains"].get(domain_label, [domain_label])
        pool = fetch_domain_pool(
            base_url, auth_header, cfg, quarter, domain_values,
            snx_release_field, deploy_plan_field, backlog_source_field,
        )
        jql_note = f'{jql_domain_clause(cfg["domain_field"], domain_values)} [classified NA+LAC JQL union]'

        for region_label in region_labels:
            records = [r for r in pool if infer_region(r) == region_label]
            tree = build_tree(records)
            release_meta, ndp_meta = enrich_metadata(base_url, auth_header, tree)
            n_rel = len([r for r in tree if r.startswith("REL-")])
            n_ndp = len({ndp for rel in tree for ndp in tree[rel] if ndp != NO_NDP})
            summary_rows.append([region_label, domain_label, quarter, n_rel, n_ndp, len(records), jql_note])
            print(f"{region_label:4} | {domain_label:10} | {quarter} -> "
                  f"{len(records)} backlog, {n_rel} release, {n_ndp} NDP", file=sys.stderr)

            base_name = f"{region_label}_{domain_label}".replace("TDS-", "")
            combo_data.append((base_name, tree, release_meta, ndp_meta))

    return combo_data, summary_rows


def save_report_workbook(combo_data, out_path, summary_rows, updated_at, quarter,
                         include_tree=False, manual_data=None):
    """Write workbook to out_path, preserving user manual columns from prior file."""
    if manual_data is None:
        manual_data = extract_manual_data(out_path) if os.path.isfile(out_path) else {}

    wb = Workbook()
    wb.remove(wb.active)

    for base_name, tree, release_meta, ndp_meta in combo_data:
        manual_by_key = manual_data.get(base_name, {})
        add_mapping_sheet(wb, base_name, tree, release_meta, ndp_meta, manual_by_key)
        if include_tree:
            add_tree_sheet(wb, f"{base_name}_tree", tree, release_meta, ndp_meta)

    add_summary_sheet(wb, summary_rows, updated_at, quarter)
    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    wb.save(out_path)
    print(f"Saved: {out_path}", file=sys.stderr)


def main():
    cfg = load_config()
    parser = argparse.ArgumentParser(description="Region x domain RELEASE->NDP->Backlog export.")
    parser.add_argument("--quarter", default=cfg.get("default_quarter", "2026Q3"))
    parser.add_argument("--regions", help="Comma list of region labels (default: all in config)")
    parser.add_argument("--domains", help="Comma list of domain-combo labels (default: all in config)")
    parser.add_argument("--tree", action="store_true", help="Also add a collapsible tree sheet per combo")
    parser.add_argument("--output-file", help="Output .xlsx path (overrides fixed Report/ folder layout)")
    args = parser.parse_args()

    region_labels = [r.strip() for r in args.regions.split(",")] if args.regions else list(cfg["regions"].keys())
    domain_labels = [d.strip() for d in args.domains.split(",")] if args.domains else list(cfg["domains"].keys())

    combo_data, summary_rows = build_report_workbook(
        args.quarter, domain_labels=domain_labels, region_labels=region_labels,
    )

    updated_at = datetime.now()
    if args.output_file:
        out_path = args.output_file
    else:
        kind = report_kind(domain_labels, cfg)
        out_path = report_output_path(kind, args.quarter)

    save_report_workbook(combo_data, out_path, summary_rows, updated_at, args.quarter,
                         include_tree=args.tree)
    print(f"\nWrote {len(summary_rows)} combo sheet(s) to: {out_path}")


if __name__ == "__main__":
    main()
