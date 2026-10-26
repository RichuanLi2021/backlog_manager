"""Checks for the spreadsheet updater, run against workbooks built in memory.

The scenario is the one the feature exists for: a report was exported a week ago and
has been worked on since — ranked, commented, highlighted — while Jira moved on. The
refresh has to bring the Jira columns forward and leave everything else alone.
"""
import sys
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

sys.path.insert(0, str(Path(__file__).resolve().parent))
from app.schemas import BacklogRow  # noqa: E402
from app.services.spreadsheet_updater import (  # noqa: E402
    LOG_SHEET,
    MARKER_HEADER,
    build_updated_workbook,
    plan_update,
)

failures: list[str] = []

HEADERS = [
    "Rank", "Release", "Release Title", "NDP", "NDP Status", "NDP Deploy Date",
    "Backlog Key", "Summary", "Issue Type", "Status", "Assignee", "Developer",
    "Reporter", "Priority", "Biz Domain", "Req Qtr", "Labels", "Backlog Source",
    "Est h", "Spent h", "UAT ETA", "PROD ETA", "Comments",
]
COL = {name: i + 1 for i, name in enumerate(HEADERS)}


def by_header(ws) -> dict:
    """Columns of a written tab by title. The standard layout moves them, so the
    input's positions say nothing about where a column ends up."""
    return {c.value: c.column for c in ws[1] if c.value}
HIGHLIGHT = PatternFill("solid", fgColor="FFFF00")


def check(condition, message):
    if not condition:
        failures.append(message)
    return condition


def row(key, **over) -> BacklogRow:
    base = dict(
        id=abs(hash(key)) % 10000, sheet="NA_GL", region="NA", domain="GL", key=key,
        release="REL-25591", release_title="Release one", ndp="NDP-7539",
        ndp_status="Open", ndp_deploy_date="N/A", summary=f"Summary for {key}",
        issue_type="Backlog", status="20-RequirementCompleted", assignee="Karl Li",
        reporter="B Person", priority="Medium", biz_domain="GL", req_qtr="2026Q3",
        labels="", backlog_source="", est_h=0.0, spent_h=0.0, archived=False,
        rank=0, developer="", uat_eta="", prod_eta="", comments="",
    )
    return BacklogRow(**{**base, **over})


def week_old_export() -> bytes:
    """What people have been working in: three backlogs, ranked, commented, highlighted."""
    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("NA_GL")
    ws.append(HEADERS)
    ws.append([
        1, "REL-25591", "Release one", "NDP-7539", "Open", "N/A",
        "CISGL-2451", "Summary for CISGL-2451", "Backlog", "20-RequirementCompleted",
        "Karl Li", "Ann Lee", "B Person", "Medium", "GL", "2026Q3", "", "", 8, 2,
        "2026-09-15", "2026-10-01", "TPM: must ship first",
    ])
    ws.append([
        2, "REL-25591", "Release one", "NDP-7539", "Open", "N/A",
        "CISGL-2466", "Summary for CISGL-2466", "Backlog", "20-RequirementCompleted",
        "Karl Li", "", "B Person", "Medium", "GL", "2026Q3", "", "", 0, 0,
        "", "", "waiting on vendor",
    ])
    ws.append([
        3, "REL-25591", "Release one", "NDP-7539", "Open", "N/A",
        "CISGL-9999", "Dropped requirement", "Backlog", "20-RequirementCompleted",
        "Karl Li", "", "B Person", "Medium", "GL", "2026Q3", "", "", 0, 0,
        "", "", "stakeholder asked for this",
    ])
    # Somebody highlighted the row they care about, in Excel, by hand.
    for column in range(1, len(HEADERS) + 1):
        ws.cell(row=2, column=column).fill = HIGHLIGHT

    about = wb.create_sheet("about")
    about.append(["Report", "GL - 2026Q3"])
    about.append(["Jira last updated", "2026-08-06 09:58:12"])

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def current_backlog() -> list[BacklogRow]:
    """A week later: statuses moved, an NDP was approved, one was raised, one dropped."""
    return [
        row("CISGL-2451", status="60-AccurateEstDone", ndp_status="QC Approved",
            ndp_deploy_date="2026-09-30", est_h=12.0, spent_h=4.0),
        row("CISGL-2466", status="50-BacklogAccepted"),
        # Raised after the file was exported, so the file has never seen it.
        row("CISGL-2556", summary="Late-breaking requirement", status="10-Placeholder"),
    ]


def test_plan_reports_the_week_of_movement():
    plan = plan_update(
        week_old_export(), "GL_2026Q3.xlsx", current_backlog(), "GL", "2026Q3", None
    )
    moved = {(c.key, c.field): c for c in plan.changes}

    check(plan.sheets[0].resolved_to == "NA_GL", f"sheet not resolved: {plan.sheets}")
    check(moved[("CISGL-2451", "status")].in_file == "20-RequirementCompleted",
          "old status not read from the file")
    check(moved[("CISGL-2451", "status")].current == "60-AccurateEstDone",
          "new status not read from the platform")
    check(moved[("CISGL-2451", "ndp_status")].current == "QC Approved", "NDP status missed")
    check(("CISGL-2451", "est_h") in moved, "hours change missed")
    check(("CISGL-2466", "status") in moved, "second status change missed")
    check(("CISGL-2466", "ndp_status") not in moved, "unchanged NDP status was flagged")

    check([n.key for n in plan.new_rows] == ["CISGL-2556"],
          f"backlog raised since the export not found: {plan.new_rows}")
    check([(g.key, g.reason) for g in plan.gone_rows] == [
        ("CISGL-9999", "not in this report any more")
    ], f"dropped backlog not flagged: {plan.gone_rows}")
    check(plan.unchanged_cells > 20, f"unchanged cells not counted: {plan.unchanged_cells}")
    check(plan.field_counts.get("status") == 2, f"field counts wrong: {plan.field_counts}")
    print(f"plan: {len(plan.changes)} cells stale, {len(plan.new_rows)} new, "
          f"{len(plan.gone_rows)} gone, {plan.unchanged_cells} already current")


def test_progress_counts_file_against_now():
    """The file has three backlogs in 20-; a week on, Jira has them spread out."""
    # One NDP has one status in a real sync, so every row under NDP-7539 agrees here.
    current = [
        row.model_copy(update={"ndp_status": "QC Approved"}) for row in current_backlog()
    ]
    plan = plan_update(week_old_export(), "GL_2026Q3.xlsx", current, "GL", "2026Q3", None)
    got = {c.status: (c.in_file, c.now) for c in plan.progress.backlogs}
    want = {
        "10-Placeholder": (0, 1),           # CISGL-2556, raised since
        "20-RequirementCompleted": (3, 0),  # all three, back then
        "50-BacklogAccepted": (0, 1),       # CISGL-2466 moved on
        "60-AccurateEstDone": (0, 1),       # CISGL-2451 moved on
    }
    check(got == want, f"backlog progress wrong: {got}")
    check([c.status for c in plan.progress.backlogs] == sorted(want),
          "statuses should run in workflow order, 10- before 20- before 50-")
    check(plan.progress.backlogs_in_file == 3 and plan.progress.backlogs_now == 3,
          f"totals wrong: {plan.progress.backlogs_in_file} / {plan.progress.backlogs_now}")
    ndps = {c.status: (c.in_file, c.now) for c in plan.progress.ndps}
    check(ndps == {"Open": (1, 0), "QC Approved": (0, 1)},
          f"NDP progress should count NDP-7539 once, Open then QC Approved now: {ndps}")
    print(f"progress: 20- 3 then 0; 10-/50-/60- one each now; NDP-7539 Open -> QC Approved")


def test_a_repeated_priority_tab_is_not_counted_twice():
    """Working files repeat NA_GL as a reordered priority-list tab: same backlogs."""
    wb = load_workbook(BytesIO(week_old_export()))
    copy = wb.copy_worksheet(wb["NA_GL"])
    copy.title = "Q4_NA_GL_Priority_List"
    buffer = BytesIO()
    wb.save(buffer)
    current = [r.model_copy(update={"ndp_status": "QC Approved"}) for r in current_backlog()]
    plan = plan_update(buffer.getvalue(), "GL.xlsx", current, "GL", "2026Q3", None)
    # Beside a priority-list tab, NA_GL is a frozen record: only the copy is worked on.
    check([(s.name, s.resolved_to) for s in plan.sheets] == [("Q4_NA_GL_Priority_List", "NA_GL")],
          f"only the priority tab should be active, read as NA_GL: {plan.sheets}")
    check(any(i.startswith("NA_GL (kept as a record") for i in plan.sheets_ignored),
          f"NA_GL should be reported as kept as a record: {plan.sheets_ignored}")
    check(plan.progress.backlogs_in_file == 3,
          f"3 backlogs, not {plan.progress.backlogs_in_file}: the copy tab was double-counted")
    check(plan.progress.backlogs_now == 3, f"now double-counted: {plan.progress.backlogs_now}")
    print("priority tab beside NA_GL: only the priority tab worked on, 3 backlogs counted")


def test_refresh_keeps_every_edit_and_the_highlight():
    updated, plan = build_updated_workbook(
        week_old_export(), "GL_2026Q3.xlsx", current_backlog(), "GL", "2026Q3", None
    )
    wb = load_workbook(BytesIO(updated))
    ws = wb["NA_GL"]
    COL = by_header(ws)  # noqa: N806
    check("IT Comment" in COL and "Comments" not in COL and "TPM Comment" in COL,
          f"Comments should have become IT Comment beside a new TPM Comment: {list(COL)}")
    by_key = {ws.cell(row=r, column=COL["Backlog Key"]).value: r for r in range(2, ws.max_row + 1)}

    first = by_key["CISGL-2451"]
    check(ws.cell(row=first, column=COL["Status"]).value == "60-AccurateEstDone",
          "status was not refreshed")
    check(ws.cell(row=first, column=COL["NDP Status"]).value == "QC Approved",
          "NDP status was not refreshed")
    check(ws.cell(row=first, column=COL["Est h"]).value == 12.0, "hours were not refreshed")

    # Everything a person put in the file has to come back untouched.
    check(ws.cell(row=first, column=COL["Rank"]).value == 1, "rank was overwritten")
    check(ws.cell(row=first, column=COL["IT Comment"]).value == "TPM: must ship first",
          "comment was overwritten")
    check(ws.cell(row=first, column=COL["Developer"]).value == "Ann Lee",
          "developer was overwritten")
    check(ws.cell(row=first, column=COL["UAT ETA"]).value == "2026-09-15",
          "UAT ETA was overwritten")
    check(ws.cell(row=first, column=COL["IT Comment"]).fill.fgColor.rgb == "00FFFF00",
          "the highlight did not survive the round trip")

    gone = by_key["CISGL-9999"]
    marker = next(c.column for c in ws[1] if c.value == MARKER_HEADER)
    check("not in this report" in (ws.cell(row=gone, column=marker).value or ""),
          "dropped backlog was not flagged in the sheet")
    check(ws.cell(row=gone, column=COL["IT Comment"]).value == "stakeholder asked for this",
          "a dropped row lost its comment")

    added = by_key.get("CISGL-2556")
    check(added is not None, "the new backlog was not appended")
    if added:
        check(ws.cell(row=added, column=COL["Summary"]).value == "Late-breaking requirement",
              "appended row has no summary")
        check(ws.cell(row=added, column=COL["Status"]).value == "10-Placeholder",
              "appended row has no status")
        check(not ws.cell(row=added, column=COL["IT Comment"]).value,
              "appended row invented a comment")
        check("new since" in (ws.cell(row=added, column=marker).value or ""),
              "appended row was not marked as new")

    log = wb[LOG_SHEET]
    entries = [r for r in log.iter_rows(values_only=True) if r[0] == "NA_GL"]
    check(len(entries) == len(plan.changes) + 1, f"log has {len(entries)} rows for "
          f"{len(plan.changes)} changes plus one addition")
    check(wb["about"].cell(row=2, column=2).value != "2026-08-06 09:58:12",
          "the about sheet still claims the old timestamp")
    print(f"refresh: {len(plan.changes)} cells rewritten, highlight and edits intact, "
          f"log has {len(entries)} entries")


def test_only_chosen_columns_are_touched():
    updated, _ = build_updated_workbook(
        week_old_export(), "GL_2026Q3.xlsx", current_backlog(), "GL", "2026Q3", None,
        fields=("status",), add_new=False, flag_gone=False,
    )
    ws = load_workbook(BytesIO(updated))["NA_GL"]
    COL = by_header(ws)  # noqa: N806
    check(ws.cell(row=2, column=COL["Status"]).value == "60-AccurateEstDone",
          "chosen column was not refreshed")
    check(ws.cell(row=2, column=COL["NDP Status"]).value == "Open",
          "an unchosen column was refreshed anyway")
    check(ws.cell(row=2, column=COL["Est h"]).value == 8, "hours changed despite not being chosen")
    keys = [ws.cell(row=r, column=COL["Backlog Key"]).value for r in range(2, ws.max_row + 1)]
    check("CISGL-2556" not in keys, "a row was appended although add_new was off")
    check(all(c.value != MARKER_HEADER for c in ws[1]), "a marker column appeared unasked")
    print(f"scoped refresh: only Status moved, {len(keys)} rows still in the sheet")


def test_a_file_for_the_wrong_report_says_so():
    plan = plan_update(
        week_old_export(), "GL_2026Q3.xlsx",
        [row("MCAM-1", sheet="LAC_AR"), row("MCAM-2", sheet="LAC_AR")],
        "AR", "2026Q3", None,
    )
    check(plan.sheets[0].resolved_to is None, "an unrelated sheet was matched anyway")
    check(any("could not be matched" in w for w in plan.warnings),
          f"no warning about the mismatch: {plan.warnings}")
    check(plan.changes == [], "changes were proposed for an unmatched sheet")
    print("wrong report: refused to touch the sheet and said why")


def main():
    for test in (
        test_plan_reports_the_week_of_movement,
        test_progress_counts_file_against_now,
        test_a_repeated_priority_tab_is_not_counted_twice,
        test_refresh_keeps_every_edit_and_the_highlight,
        test_only_chosen_columns_are_touched,
        test_a_file_for_the_wrong_report_says_so,
    ):
        try:
            test()
        except Exception as exc:  # a crash is a failure, not a stack trace
            failures.append(f"{test.__name__} raised {type(exc).__name__}: {exc}")

    if failures:
        print("\nFAILURES:")
        for f in failures:
            print(f"  - {f}")
        return 1
    print("\nPASS: a week-old file comes forward on Jira columns and keeps every edit")
    return 0


if __name__ == "__main__":
    sys.exit(main())
