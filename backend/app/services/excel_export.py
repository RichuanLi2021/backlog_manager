"""Write a report view to an .xlsx workbook, one worksheet per region sheet."""

from __future__ import annotations

from datetime import datetime
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from ..schemas import BacklogRow

# The team's standard layout: NA_GL's columns, with Package after Assignee and the
# old Comments split into IT Comment and TPM Comment. Every exported tab and every
# priority-list tab the Updater rewrites follows it, and the on-screen grid too.
COLUMNS: list[tuple[str, str]] = [
    ("rank", "Rank"),
    ("release", "Release"),
    ("release_title", "Release Title"),
    ("ndp", "NDP"),
    ("ndp_status", "NDP Status"),
    ("uat_eta", "UAT ETA"),
    ("prod_eta", "PROD ETA"),
    ("ndp_deploy_date", "NDP Deploy Date"),
    ("key", "Backlog Key"),
    ("summary", "Summary"),
    ("issue_type", "Issue Type"),
    ("status", "Status"),
    ("assignee", "Assignee"),
    ("package", "Package"),
    ("developer", "Developer"),
    ("reporter", "Reporter"),
    ("priority", "Priority"),
    ("biz_domain", "Biz Domain"),
    ("req_qtr", "Req Qtr"),
    ("labels", "Labels"),
    ("backlog_source", "Backlog Source"),
    ("est_h", "Est h"),
    ("spent_h", "Spent h"),
    ("comments", "IT Comment"),
    ("tpm_comment", "TPM Comment"),
]

WIDE = {"release_title", "summary", "comments", "tpm_comment"}
NUMERIC = {"rank", "est_h", "spent_h"}

HDR_FILL = PatternFill("solid", fgColor="4472C4")
WHITE_BOLD = Font(bold=True, color="FFFFFF")

# Same palette the skill uses for NDP status cells.
NDP_FILLS = {
    "QC Approved": PatternFill("solid", fgColor="F8CBAD"),
    "IIC Approved": PatternFill("solid", fgColor="FFF2CC"),
    "IT Exec Approved": PatternFill("solid", fgColor="FFE699"),
    "Verified": PatternFill("solid", fgColor="C6EFCE"),
}


def _width(field: str) -> int:
    if field in WIDE:
        return 55
    if field in NUMERIC:
        return 8
    return 18


NDP_STATUS_COL = [field for field, _ in COLUMNS].index("ndp_status") + 1


def priority_tab(sheet: str, quarter: str) -> str:
    """'NA_GL' in 2026Q4 -> 'Q4_NA_GL_Priority_List', the tab the team works in."""
    return f"{quarter[-2:]}_{sheet}_Priority_List"[:31]


def build_workbook(
    rows: list[BacklogRow], about: list[tuple[str, str]], quarter: str = ""
) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)

    for sheet in dict.fromkeys(r.sheet for r in rows):
        ws = wb.create_sheet(priority_tab(sheet, quarter) if quarter else sheet[:31])
        ws.append([header for _, header in COLUMNS])
        for cell in ws[1]:
            cell.fill = HDR_FILL
            cell.font = WHITE_BOLD

        for row in (r for r in rows if r.sheet == sheet):
            ws.append([getattr(row, field) for field, _ in COLUMNS])
            fill = NDP_FILLS.get((row.ndp_status or "").strip())
            if fill:
                ws.cell(row=ws.max_row, column=NDP_STATUS_COL).fill = fill

        for index, (field, _) in enumerate(COLUMNS, start=1):
            ws.column_dimensions[get_column_letter(index)].width = _width(field)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNS))}{ws.max_row}"

    info = wb.create_sheet("about")
    for label, value in [*about, ("Exported", datetime.now().strftime("%Y-%m-%d %H:%M:%S"))]:
        info.append([label, value])
        info.cell(row=info.max_row, column=1).font = Font(bold=True)
        info.cell(row=info.max_row, column=2).alignment = Alignment(wrap_text=True)
    info.column_dimensions["A"].width = 22
    info.column_dimensions["B"].width = 70

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
