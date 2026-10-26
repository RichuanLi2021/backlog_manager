"""Put a working tab into the team's standard layout without losing anything in it.

The standard is `excel_export.COLUMNS`: NA_GL's columns, Package after Assignee, and
the old Comments split into IT Comment and TPM Comment. Circulated tabs drift from
it — columns reordered, Package inserted, Comments still in one piece — so before a
tab is updated it is rearranged into that order:

- every recognised column moves to its standard place under its standard title,
  so a legacy "Comments" column becomes "IT Comment", text and all;
- a standard column the tab lacks is added empty (TPM Comment, the first time);
- anything unrecognised is kept, after the standard columns, never dropped;
- "Update Status", the Updater's own note column, stays last.

Values, fills, fonts, number formats and widths move with their column. Running it on
a tab already in the standard layout changes nothing.
"""
from __future__ import annotations

from copy import copy
from io import BytesIO

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .excel_export import COLUMNS
from .spreadsheet_layout import (
    ALL_ALIASES,
    JIRA_FIELDS,
    USER_FIELDS,
    SheetLayout,
    is_priority_tab,
    match_field,
    read_layouts,
    text_of,
)

MARKER = "update status"
WIDE_FIELDS = {"release_title", "summary", "comments", "tpm_comment"}


def active_tabs(sheets: list[SheetLayout]) -> list[SheetLayout]:
    """The tabs people work in. Once a file has priority-list tabs, those are the
    working copies and the plain region tabs are kept only as a frozen record."""
    priority = [s for s in sheets if is_priority_tab(s.name)]
    return priority or sheets


def frozen_tabs(sheets: list[SheetLayout]) -> list[str]:
    working = {s.name for s in active_tabs(sheets)}
    return [s.name for s in sheets if s.name not in working]


def _column_empty(sheet: SheetLayout, column: int) -> bool:
    return all(
        column > len(row) or row[column - 1] in (None, "") for row in sheet.grid
    )


def plan_order(sheet: SheetLayout) -> tuple[list[int | None], list[str], list[str]]:
    """The standard arrangement of one tab.

    Returns the old column number to place in each new position (None for a column
    to add), the title each position gets, and plain-language notes on what moves.
    """
    header = sheet.grid[sheet.header_row - 1] if sheet.grid else ()
    width = max((len(r) for r in sheet.grid), default=0)
    found: dict[str, int] = {}
    for column in range(1, width + 1):
        value = header[column - 1] if column <= len(header) else None
        field = match_field(value, ALL_ALIASES)
        if field and field not in found:
            found[field] = column

    order: list[int | None] = []
    titles: list[str] = []
    notes: list[str] = []
    used: set[int] = set()
    for field, title in COLUMNS:
        column = found.get(field)
        order.append(column)
        titles.append(title)
        if column is None:
            notes.append(f"adds {title}")
            continue
        used.add(column)
        was = text_of(header[column - 1]) if column <= len(header) else ""
        if was != title:
            notes.append(f'"{was}" becomes "{title}"')

    marker = None
    for column in range(1, width + 1):
        if column in used:
            continue
        value = text_of(header[column - 1]) if column <= len(header) else ""
        if value.strip().lower() == MARKER:
            marker = column
        elif value or not _column_empty(sheet, column):
            order.append(column)  # unrecognised but holds something: keep it
            titles.append(value)
            notes.append(f'keeps "{value or get_column_letter(column)}" after the standard columns')
    if marker:
        order.append(marker)
        titles.append("Update Status")

    kept = [c for c in order if c is not None]
    moved = kept != sorted(kept)
    if moved:
        notes.append("reorders the columns")
    # Keeping an extra column is only news when the tab is being rearranged anyway.
    changes = moved or None in order or any(" becomes " in n for n in notes)
    return order, titles, notes if changes else []


def describe(data: bytes, filename: str) -> list[str]:
    """What standardising would do, one line per tab that needs it."""
    layout = read_layouts(data, filename, JIRA_FIELDS + USER_FIELDS, ALL_ALIASES)
    lines = []
    for sheet in active_tabs(layout.sheets):
        _, _, notes = plan_order(sheet)
        if notes:
            lines.append(f"{sheet.name}: " + "; ".join(notes))
    return lines


def standardize(data: bytes, filename: str) -> tuple[bytes, list[str]]:
    """The workbook with every working tab in the standard layout, and what changed."""
    layout = read_layouts(data, filename, JIRA_FIELDS + USER_FIELDS, ALL_ALIASES)
    plans = {
        sheet.name: (sheet, *plan_order(sheet)) for sheet in active_tabs(layout.sheets)
    }
    todo = {name: p for name, p in plans.items() if p[3]}
    if not todo:
        return data, []

    wb = load_workbook(BytesIO(data))
    lines = []
    for name, (sheet, order, titles, notes) in todo.items():
        _rearrange(wb[name], sheet.header_row, order, titles)
        lines.append(f"{name}: " + "; ".join(notes))
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue(), lines


def _rearrange(ws, header_row: int, order: list[int | None], titles: list[str]):
    last_row = ws.max_row
    old_width = ws.max_column

    # Take everything first: positions are about to be overwritten in place.
    cells = {}
    widths = {}
    for column in {c for c in order if c is not None}:
        letter = get_column_letter(column)
        widths[column] = (ws.column_dimensions[letter].width, ws.column_dimensions[letter].hidden)
        cells[column] = [
            (
                ws.cell(row=r, column=column).value,
                copy(ws.cell(row=r, column=column)._style),
                ws.cell(row=r, column=column).hyperlink,
                copy(ws.cell(row=r, column=column).comment) if ws.cell(row=r, column=column).comment else None,
            )
            for r in range(1, last_row + 1)
        ]

    fields = [f for f, _ in COLUMNS]
    for position, column in enumerate(order, start=1):
        letter = get_column_letter(position)
        if column is not None:
            for r, (value, style, link, note) in enumerate(cells[column], start=1):
                cell = ws.cell(row=r, column=position)
                cell.value = value
                cell._style = style
                cell.hyperlink = link.target if link is not None and link.target else None
                cell.comment = note
            width, hidden = widths[column]
            ws.column_dimensions[letter].width = width
            ws.column_dimensions[letter].hidden = hidden
        else:
            # A new column borrows the look of its left neighbour, row by row, so a
            # highlighted row stays highlighted all the way across.
            for r in range(1, last_row + 1):
                cell = ws.cell(row=r, column=position)
                cell.value = None
                cell.comment = None
                if position > 1:
                    cell._style = copy(ws.cell(row=r, column=position - 1)._style)
            field = fields[position - 1] if position <= len(fields) else ""
            ws.column_dimensions[letter].width = 55 if field in WIDE_FIELDS else 18
            ws.column_dimensions[letter].hidden = False
        if position <= len(titles) and titles[position - 1]:
            ws.cell(row=header_row, column=position).value = titles[position - 1]

    # Anything past the new right edge was moved left; clear the leftovers.
    for column in range(len(order) + 1, old_width + 1):
        for r in range(1, last_row + 1):
            cell = ws.cell(row=r, column=column)
            cell.value = None
            cell.comment = None
            cell._style = copy(ws.cell(row=r, column=old_width + 1)._style)

    _move_filters(ws, header_row, order, last_row)


def _move_filters(ws, header_row: int, order: list[int | None], last_row: int):
    """Keep each filter on the column it was set on, wherever that column moved.

    A filter is stored as an offset from the filter range's first column, not by
    header. Left alone after a reorder it silently re-points at whatever column now
    sits at that offset: an ETA filter ends up on Reporter and hides most rows.
    """
    auto = ws.auto_filter
    if not auto.ref:
        return
    from openpyxl.utils.cell import range_boundaries

    first_old = range_boundaries(auto.ref)[0]
    new_position = {old: new for new, old in enumerate(order, start=1) if old is not None}
    kept = []
    for criteria in auto.filterColumn:
        moved_to = new_position.get(first_old + criteria.colId)
        if moved_to is not None:  # a filter on a column that no longer exists is dropped
            criteria.colId = moved_to - 1
            kept.append(criteria)
    auto.filterColumn = kept
    # The sort indicator names cell ranges by the old positions; the rows are already
    # in that order, so dropping the indicator loses nothing.
    auto.sortState = None
    auto.ref = f"A{header_row}:{get_column_letter(len(order))}{last_row}"
